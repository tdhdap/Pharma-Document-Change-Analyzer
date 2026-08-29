import os
import tempfile
import zipfile

from docx import Document as DocxDocument
from lxml import etree
from docx.oxml.ns import qn

from app import pipeline, llm_classifier
from app.extraction import extract_text
from app.models import Paragraph, LLMClassification, TableCoordinate


_TEST_DOCS_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "test-documents",
    "docx",
)


def test_pipeline_reproduces_the_outline_example(monkeypatch):
    old_paragraphs = [
        Paragraph(text="Assay acceptance criterion: 95.0% to 105.0%."),
        Paragraph(text="Samples shall be stored at 25°C ± 2°C."),
        Paragraph(text="The Quality Control Manager shall approve the result."),
    ]
    new_paragraphs = [
        Paragraph(text="Assay acceptance criterion: 98.0% to 102.0%."),
        Paragraph(text="Samples shall be stored at 25°C ± 2°C and 60% RH ± 5% RH."),
        Paragraph(text="The Quality Assurance Manager shall approve the result."),
    ]

    def fake_classify(unresolved):
        results = []
        for item in unresolved:
            if "Manager" in item["old_text"]:
                results.append(LLMClassification(
                    change_id=item["change_id"],
                    change_type="role_responsibility_change",
                    reason="Approval responsibility changed from QC Manager to QA Manager.",
                    confidence=0.9,
                ))
            else:
                results.append(LLMClassification(
                    change_id=item["change_id"],
                    change_type="qualitative_specification_change",
                    reason="An additional environmental control requirement (humidity) was added.",
                    confidence=0.85,
                ))
        return results

    monkeypatch.setattr(llm_classifier, "classify_changes_batch", fake_classify)

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "SOP_v1.txt", "SOP_v2.txt")

    # 5 rows total: line 1 -> numeric_change (1); line 2 -> unit_change + numeric_change
    # + one AI-classified row for the added humidity clause (3); line 3 -> one AI-classified
    # role_responsibility_change row (1).
    assert result.summary.total_changes == 5
    assert result.summary.high_risk == 4
    assert result.summary.medium_risk == 1

    change_types = {c.change_type for c in result.changes}
    assert "numeric_change" in change_types
    assert "unit_change" in change_types
    assert "role_responsibility_change" in change_types

    # This is the row this whole feature exists to add: line 2 alone must now produce
    # 3 distinct rows instead of 1.
    samples_changes = [c for c in result.changes if c.old_text == "Samples shall be stored at 25°C ± 2°C."]
    assert len(samples_changes) == 3
    assert {c.change_type for c in samples_changes} == {
        "unit_change", "numeric_change", "qualitative_specification_change",
    }

    role_change = next(c for c in result.changes if c.change_type == "role_responsibility_change")
    assert role_change.ai_risk_level == "Medium"
    assert "QA Manager" in role_change.reason


def test_two_regex_detections_on_one_line_produce_two_rows_and_skip_ai(monkeypatch):
    old_paragraphs = [Paragraph(text="Dispense 10 mL of solution.")]
    new_paragraphs = [Paragraph(text="Dispense 20 L of solution.")]

    def fail_if_called(unresolved):
        raise AssertionError("classify_changes_batch should not be called - regex fully explains this line")

    monkeypatch.setattr(llm_classifier, "classify_changes_batch", fail_if_called)

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.txt", "new.txt")

    assert result.summary.total_changes == 2
    assert {c.change_type for c in result.changes} == {"unit_change", "numeric_change"}


def test_regex_detection_plus_separate_wording_change_produces_two_rows(monkeypatch):
    old_paragraphs = [Paragraph(text="Weigh 50 mg of sample, thoroughly mixed.")]
    new_paragraphs = [Paragraph(text="Weigh 55 mg of sample, completely mixed.")]

    def fake_classify(unresolved):
        assert len(unresolved) == 1
        assert unresolved[0]["already_detected"] == ["numeric_change"]
        return [LLMClassification(
            change_id=unresolved[0]["change_id"],
            change_type="qualitative_specification_change",
            reason="The mixing descriptor was changed from thoroughly to completely.",
            confidence=0.9,
        )]

    monkeypatch.setattr(llm_classifier, "classify_changes_batch", fake_classify)

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.txt", "new.txt")

    assert result.summary.total_changes == 2
    change_types = {c.change_type for c in result.changes}
    assert change_types == {"numeric_change", "qualitative_specification_change"}


def test_date_and_numeric_change_on_one_line_both_survive_with_independent_risk(monkeypatch):
    old_paragraphs = [Paragraph(text="Sample collected on 01 Jan 2024, weight 50 mg.")]
    new_paragraphs = [Paragraph(text="Sample collected on 15 Mar 2024, weight 55 mg.")]

    def fail_if_called(unresolved):
        raise AssertionError("classify_changes_batch should not be called - regex fully explains this line")

    monkeypatch.setattr(llm_classifier, "classify_changes_batch", fail_if_called)

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.txt", "new.txt")

    assert result.summary.total_changes == 2
    risks_by_type = {c.change_type: c.ai_risk_level for c in result.changes}
    assert risks_by_type == {"date_change": "Medium", "numeric_change": "High"}


def test_pending_change_never_leaks_when_llm_response_omits_it(monkeypatch):
    old_paragraphs = [
        Paragraph(text="The Quality Control Manager shall approve the result."),
    ]
    new_paragraphs = [
        Paragraph(text="The Quality Assurance Manager shall approve the result."),
    ]

    def fake_classify_returns_nothing(unresolved):
        return []  # simulates a partial/malformed Gemini response that omits this item

    monkeypatch.setattr(llm_classifier, "classify_changes_batch", fake_classify_returns_nothing)

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "SOP_v1.txt", "SOP_v2.txt")

    assert result.summary.total_changes == 1
    change = result.changes[0]
    assert change.change_type == "unclassified"
    assert change.change_type != "pending_llm_classification"
    assert change.ai_risk_level == "Medium"
    assert change.confidence == 0.0


def test_table_cell_deletion_gets_table_labeled_change_type():
    old_paragraphs = [Paragraph(text="Old Row Value", from_table=True)]
    new_paragraphs = []

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.docx", "new.docx")

    assert len(result.changes) == 1
    change = result.changes[0]
    assert change.source == "Table"
    assert change.change_type == "deleted_table_content"
    assert change.reason == "Table content removed."


def test_table_cell_addition_gets_table_labeled_change_type():
    old_paragraphs = []
    new_paragraphs = [Paragraph(text="New Row Value", from_table=True)]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.docx", "new.docx")

    assert len(result.changes) == 1
    change = result.changes[0]
    assert change.source == "Table"
    assert change.change_type == "added_table_content"
    assert change.reason == "New table content added."


def test_body_paragraph_addition_keeps_generic_change_type():
    old_paragraphs = []
    new_paragraphs = [Paragraph(text="New body sentence.")]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.docx", "new.docx")

    assert len(result.changes) == 1
    change = result.changes[0]
    assert change.source == "Body"
    assert change.change_type == "added_paragraph"
    assert change.reason == "New paragraph added."


def test_table_cell_with_precise_regex_change_keeps_precise_type():
    old_paragraphs = [Paragraph(text="Weigh 10 mg of sample.", from_table=True)]
    new_paragraphs = [Paragraph(text="Weigh 20 mg of sample.", from_table=True)]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.docx", "new.docx")

    assert len(result.changes) == 1
    change = result.changes[0]
    assert change.source == "Table"
    assert change.change_type == "numeric_change"


def test_pipeline_detects_renumbering_and_reordering_together():
    old_paragraphs = [
        Paragraph(text="1.0 Scope"),
        Paragraph(text="This procedure applies to testing performed in the QC lab."),
        Paragraph(text="2.0 Acceptance Criteria"),
        Paragraph(text="Results shall conform to the specified limits."),
        Paragraph(text="3.0 Approval"),
        Paragraph(text="The Quality Assurance Manager shall approve results."),
    ]
    new_paragraphs = [
        Paragraph(text="1.0 Scope"),
        Paragraph(text="This procedure applies to testing performed in the QC lab."),
        Paragraph(text="2.0 Approval"),
        Paragraph(text="The Quality Assurance Manager shall approve results."),
        Paragraph(text="3.0 Acceptance Criteria"),
        Paragraph(text="Results shall conform to the specified limits."),
    ]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "SOP_v1.txt", "SOP_v2.txt")

    renumbered = {c.old_text: c for c in result.changes if c.change_type == "section_renumbered"}
    reordered = [c for c in result.changes if c.change_type == "section_reordered"]

    assert set(renumbered.keys()) == {"2.0 Acceptance Criteria", "3.0 Approval"}
    assert renumbered["2.0 Acceptance Criteria"].new_text == "3.0 Acceptance Criteria"
    assert renumbered["2.0 Acceptance Criteria"].reason == "Section renumbered from '2.0' to '3.0'."
    assert renumbered["3.0 Approval"].new_text == "2.0 Approval"
    assert renumbered["3.0 Approval"].reason == "Section renumbered from '3.0' to '2.0'."
    for c in renumbered.values():
        assert c.ai_risk_level == "Informational"
        assert c.source == "Body"

    assert len(reordered) == 1
    assert reordered[0].section == "3.0 Approval"
    assert reordered[0].old_text == "3.0 Approval"
    assert reordered[0].new_text == "2.0 Approval"
    assert reordered[0].reason == "Section moved from position 3 to position 2 in the document."
    assert reordered[0].ai_risk_level == "Informational"


def test_pipeline_reports_a_new_section_as_section_added_not_scattered_paragraphs():
    old_paragraphs = [
        Paragraph(text="1.0 Scope"),
        Paragraph(text="This procedure applies to testing performed in the QC lab."),
    ]
    new_paragraphs = [
        Paragraph(text="1.0 Scope"),
        Paragraph(text="This procedure applies to testing performed in the QC lab."),
        Paragraph(text="2.0 Environmental Monitoring"),
        Paragraph(text="Environmental monitoring shall be performed weekly using settle plates."),
    ]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "SOP_v1.txt", "SOP_v2.txt")

    assert result.summary.total_changes == 1
    change = result.changes[0]
    assert change.change_type == "section_added"
    assert change.section == "2.0 Environmental Monitoring"
    assert change.new_text == (
        "2.0 Environmental Monitoring\n"
        "Environmental monitoring shall be performed weekly using settle plates."
    )
    assert change.ai_risk_level == "High"


def test_pipeline_reports_a_deleted_section_as_section_deleted_not_scattered_paragraphs():
    old_paragraphs = [
        Paragraph(text="1.0 Scope"),
        Paragraph(text="This procedure applies to testing performed in the QC lab."),
        Paragraph(text="8.0 Deviation Handling"),
        Paragraph(text="Any deviation from this procedure shall be documented and approved."),
    ]
    new_paragraphs = [
        Paragraph(text="1.0 Scope"),
        Paragraph(text="This procedure applies to testing performed in the QC lab."),
    ]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "SOP_v1.txt", "SOP_v2.txt")

    assert result.summary.total_changes == 1
    change = result.changes[0]
    assert change.change_type == "section_deleted"
    assert change.section == "8.0 Deviation Handling"
    assert change.old_text == (
        "8.0 Deviation Handling\n"
        "Any deviation from this procedure shall be documented and approved."
    )
    assert change.ai_risk_level == "High"


def test_pipeline_headingless_paragraph_addition_still_uses_added_paragraph():
    # Regression guard: sectioning.py wraps a headingless document's paragraphs in
    # synthetic "Paragraph N" sections. Without the synthetic-heading guard, this
    # would get misreported as a High-risk section_added instead of the existing
    # Medium-risk added_paragraph. This mirrors the pre-existing
    # test_body_paragraph_addition_keeps_generic_change_type case but asserts the
    # change_type explicitly against this plan's new code path.
    old_paragraphs = []
    new_paragraphs = [Paragraph(text="New body sentence with no heading anywhere.")]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.txt", "new.txt")

    assert result.summary.total_changes == 1
    change = result.changes[0]
    assert change.change_type == "added_paragraph"
    assert change.ai_risk_level != "High"


def test_pipeline_content_moved_into_a_new_section_is_reported_as_moved_not_deleted():
    # Regression test for the final-review finding: content relocating from a
    # matched section into a brand-new section must still produce moved_paragraph,
    # not a false deleted_paragraph plus a duplicated line inside section_added.
    old_paragraphs = [
        Paragraph(text="1.0 Materials"),
        Paragraph(text="All raw materials shall be inspected upon receipt for identity and quality."),
        Paragraph(text="Material shall not be used beyond 24 months from the date of manufacture."),
        Paragraph(text="2.0 Approval"),
        Paragraph(text="The Quality Assurance Manager shall approve results."),
    ]
    new_paragraphs = [
        Paragraph(text="1.0 Materials"),
        Paragraph(text="All raw materials shall be inspected upon receipt for identity and quality."),
        Paragraph(text="2.0 Approval"),
        Paragraph(text="The Quality Assurance Manager shall approve results."),
        Paragraph(text="3.0 Storage and Shelf Life"),
        Paragraph(text="Material shall not be used beyond 24 months from the date of manufacture."),
    ]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.txt", "new.txt")

    change_types = [c.change_type for c in result.changes]
    assert "deleted_paragraph" not in change_types
    moved = [c for c in result.changes if c.change_type == "moved_paragraph"]
    assert len(moved) == 1
    assert moved[0].old_text == "Material shall not be used beyond 24 months from the date of manufacture."
    section_added = [c for c in result.changes if c.change_type == "section_added"]
    assert len(section_added) == 1
    assert section_added[0].new_text == "3.0 Storage and Shelf Life"


def test_pipeline_reports_both_renumbering_and_heading_change_for_one_section(monkeypatch):
    old_paragraphs = [
        Paragraph(text="1.0 Scope"),
        Paragraph(text="This procedure applies to testing performed in the QC lab."),
        Paragraph(text="8.0 Deviation Handling"),
        Paragraph(text="Any deviation shall be documented within 24 hours and approved."),
    ]
    new_paragraphs = [
        Paragraph(text="1.0 Scope"),
        Paragraph(text="This procedure applies to testing performed in the QC lab."),
        Paragraph(text="9.0 Non-Conformance Management"),
        Paragraph(text="Any deviation shall be documented within 48 hours and approved."),
    ]

    def fail_if_called(unresolved):
        raise AssertionError("classify_changes_batch should not be called - regex fully explains this line")

    monkeypatch.setattr(llm_classifier, "classify_changes_batch", fail_if_called)

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "SOP_v1.txt", "SOP_v2.txt")

    renumbered = [c for c in result.changes if c.change_type == "section_renumbered"]
    heading_changed = [c for c in result.changes if c.change_type == "section_heading_changed"]
    numeric_changed = [c for c in result.changes if c.change_type == "numeric_change"]

    assert len(renumbered) == 1
    assert renumbered[0].section == "8.0 Deviation Handling"
    assert renumbered[0].reason == "Section renumbered from '8.0' to '9.0'."

    assert len(heading_changed) == 1
    assert heading_changed[0].section == "8.0 Deviation Handling"
    assert heading_changed[0].old_text == "8.0 Deviation Handling"
    assert heading_changed[0].new_text == "9.0 Non-Conformance Management"
    assert heading_changed[0].ai_risk_level == "Medium"

    assert len(numeric_changed) == 1
    assert numeric_changed[0].ai_risk_level == "High"


def test_pipeline_reports_only_heading_changed_when_number_is_unchanged():
    old_paragraphs = [
        Paragraph(text="1.0 Scope"),
        Paragraph(text="This procedure applies to testing performed in the QC lab."),
    ]
    new_paragraphs = [
        Paragraph(text="1.0 Purpose"),
        Paragraph(text="This procedure applies to testing performed in the QC lab."),
    ]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "SOP_v1.txt", "SOP_v2.txt")

    assert result.summary.total_changes == 1
    change = result.changes[0]
    assert change.change_type == "section_heading_changed"
    assert change.old_text == "1.0 Scope"
    assert change.new_text == "1.0 Purpose"


def test_table_cell_edit_preserves_table_position_on_both_sides():
    position = TableCoordinate(table_id=0, row=1, col=1)
    old_paragraphs = [Paragraph(text="Weigh 10 mg of sample.", from_table=True, table_position=position)]
    new_paragraphs = [Paragraph(text="Weigh 20 mg of sample.", from_table=True, table_position=position)]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.docx", "new.docx")

    assert len(result.changes) == 1
    change = result.changes[0]
    assert change.old_table_position == position
    assert change.new_table_position == position


def test_table_row_addition_only_sets_new_table_position():
    position = TableCoordinate(table_id=0, row=3, col=0)
    old_paragraphs = []
    new_paragraphs = [Paragraph(text="New Row Value", from_table=True, table_position=position)]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.docx", "new.docx")

    assert len(result.changes) == 1
    change = result.changes[0]
    assert change.old_table_position is None
    assert change.new_table_position == position


def test_table_row_deletion_only_sets_old_table_position():
    position = TableCoordinate(table_id=0, row=3, col=0)
    old_paragraphs = [Paragraph(text="Old Row Value", from_table=True, table_position=position)]
    new_paragraphs = []

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.docx", "new.docx")

    assert len(result.changes) == 1
    change = result.changes[0]
    assert change.old_table_position == position
    assert change.new_table_position is None


def test_body_paragraph_addition_has_no_table_position():
    old_paragraphs = []
    new_paragraphs = [Paragraph(text="New body sentence.")]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.docx", "new.docx")

    assert len(result.changes) == 1
    change = result.changes[0]
    assert change.old_table_position is None
    assert change.new_table_position is None


def test_pipeline_adding_a_section_break_does_not_falsely_flag_header_footer_as_changed():
    # Regression test for the exact scenario the final whole-branch review found broken:
    # a document revised to add a section break, with header/footer content that doesn't
    # actually change for a reader (second section stays linked/inherited). Before the
    # fix, this spuriously produced section_heading_changed findings like
    # "Page Header" -> "Page Header (Section 1)".
    with tempfile.TemporaryDirectory() as tmp_dir:
        old_path = os.path.join(tmp_dir, "old.docx")
        new_path = os.path.join(tmp_dir, "new.docx")

        old_doc = DocxDocument()
        old_doc.add_paragraph("Body content here.")
        old_doc.sections[0].header.paragraphs[0].text = "Confidential - SOP-1234"
        old_doc.sections[0].footer.paragraphs[0].text = "Uncontrolled Copy - Page 1"
        old_doc.save(old_path)

        new_doc = DocxDocument()
        new_doc.add_paragraph("Body content here.")
        new_doc.sections[0].header.paragraphs[0].text = "Confidential - SOP-1234"
        new_doc.sections[0].footer.paragraphs[0].text = "Uncontrolled Copy - Page 1"
        new_doc.add_section()
        new_doc.add_paragraph("More content after a section break.")
        new_doc.save(new_path)

        old_extracted = extract_text(old_path, "docx")
        new_extracted = extract_text(new_path, "docx")

    result = pipeline.compare_documents(old_extracted, new_extracted, "old.docx", "new.docx")

    heading_changed = [c for c in result.changes if c.change_type == "section_heading_changed"]
    assert heading_changed == []


_TEXT_BOX_DRAWING_XML = """<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
    xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
    xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
  <wp:inline>
    <a:graphic>
      <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
        <wps:wsp>
          <wps:txbx>
            <w:txbxContent><w:p><w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:p></w:txbxContent>
          </wps:txbx>
        </wps:wsp>
      </a:graphicData>
    </a:graphic>
  </wp:inline>
</w:drawing>"""


def _add_text_box_paragraph(doc, text):
    body = doc.element.body
    p = body.makeelement(qn("w:p"), {})
    r = p.makeelement(qn("w:r"), {})
    p.append(r)
    drawing = etree.fromstring(_TEXT_BOX_DRAWING_XML.format(text=text).encode())
    r.append(drawing)
    body.append(p)


def test_pipeline_adding_an_earlier_text_box_does_not_falsely_flag_later_one_as_changed():
    # Regression test for the exact scenario this plan's final whole-branch
    # review found broken: adding an earlier text box shifts every later text
    # box's "Text Box N" number, even though its own content never changed.
    # Before the fix, this spuriously produced section_heading_changed findings
    # like "Text Box 1" -> "Text Box 2".
    with tempfile.TemporaryDirectory() as tmp_dir:
        old_path = os.path.join(tmp_dir, "old.docx")
        new_path = os.path.join(tmp_dir, "new.docx")

        old_doc = DocxDocument()
        old_doc.add_paragraph("Body content here.")
        _add_text_box_paragraph(old_doc, "Store samples at 25 C.")
        old_doc.save(old_path)

        new_doc = DocxDocument()
        new_doc.add_paragraph("Body content here.")
        _add_text_box_paragraph(new_doc, "NOTE: revision history updated.")
        _add_text_box_paragraph(new_doc, "Store samples at 25 C.")
        new_doc.save(new_path)

        old_paragraphs = extract_text(old_path, "docx")
        new_paragraphs = extract_text(new_path, "docx")

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.docx", "new.docx")

    heading_changed = [c for c in result.changes if c.change_type == "section_heading_changed"]
    assert heading_changed == []


def _add_footnote_reference(paragraph, footnote_id):
    p = paragraph._p
    r = p.makeelement(qn("w:r"), {})
    ref = r.makeelement(qn("w:footnoteReference"), {qn("w:id"): str(footnote_id)})
    r.append(ref)
    p.append(r)


def _save_docx_with_footnotes(doc, path, footnotes):
    doc.save(path)
    footnote_blocks = "".join(
        f'<w:footnote w:id="{fid}">'
        + "".join(f'<w:p><w:r><w:t xml:space="preserve">{t}</w:t></w:r></w:p>' for t in texts)
        + "</w:footnote>"
        for fid, texts in footnotes
    )
    footnotes_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="0">'
        '<w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
        f'{footnote_blocks}'
        '</w:footnotes>'
    )
    with zipfile.ZipFile(path, "r") as zin:
        names = zin.namelist()
        content_types_xml = zin.read("[Content_Types].xml").decode("utf-8")
        rels_xml = zin.read("word/_rels/document.xml.rels").decode("utf-8")
        other = {n: zin.read(n) for n in names if n not in ("[Content_Types].xml", "word/_rels/document.xml.rels")}
    content_types_xml = content_types_xml.replace(
        "</Types>",
        '<Override PartName="/word/footnotes.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>',
    )
    rels_xml = rels_xml.replace(
        "</Relationships>",
        '<Relationship Id="rIdFootnotes" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" '
        'Target="footnotes.xml"/></Relationships>',
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        zout.writestr("[Content_Types].xml", content_types_xml)
        zout.writestr("word/_rels/document.xml.rels", rels_xml)
        zout.writestr("word/footnotes.xml", footnotes_xml)
        for name, data in other.items():
            zout.writestr(name, data)


def test_pipeline_adding_an_earlier_footnote_does_not_falsely_flag_later_one_as_changed():
    # Regression test mirroring test_pipeline_adding_an_earlier_text_box_does_not_falsely_flag_later_one_as_changed:
    # adding an earlier footnote reference shifts every later footnote's "Footnote N"
    # number, even though its own content never changed. Without the guard added in
    # this plan, this would spuriously produce section_heading_changed findings like
    # "Footnote 1" -> "Footnote 2".
    with tempfile.TemporaryDirectory() as tmp_dir:
        old_path = os.path.join(tmp_dir, "old.docx")
        new_path = os.path.join(tmp_dir, "new.docx")

        old_doc = DocxDocument()
        old_p = old_doc.add_paragraph("Store samples per the stability protocol.")
        _add_footnote_reference(old_p, "1")
        _save_docx_with_footnotes(old_doc, old_path, [("1", ["Store samples at 25 C."])])

        new_doc = DocxDocument()
        new_p1 = new_doc.add_paragraph("Revision history updated for this release.")
        _add_footnote_reference(new_p1, "1")
        new_p2 = new_doc.add_paragraph("Store samples per the stability protocol.")
        _add_footnote_reference(new_p2, "2")
        _save_docx_with_footnotes(new_doc, new_path, [
            ("1", ["NOTE: revision history updated."]),
            ("2", ["Store samples at 25 C."]),
        ])

        old_paragraphs = extract_text(old_path, "docx")
        new_paragraphs = extract_text(new_path, "docx")

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.docx", "new.docx")

    heading_changed = [c for c in result.changes if c.change_type == "section_heading_changed"]
    assert heading_changed == []


def test_moved_paragraph_that_was_also_edited_reports_both_move_and_content_change(monkeypatch):
    # A relocation must never mask an edit: a spec value changing while the
    # paragraph moves is exactly the high-consequence case the 0.85 move-matching
    # threshold is most likely to swallow.
    def fail_if_called(unresolved):
        raise AssertionError("no AI call expected - the numeric regex fully explains this change")

    monkeypatch.setattr(llm_classifier, "classify_changes_batch", fail_if_called)

    old_paragraphs = [
        Paragraph(text="1.0 Purpose"),
        Paragraph(text="This SOP defines tablet compression for all product lines."),
        Paragraph(text="Compression force shall be maintained at 15 kN."),
        Paragraph(text="2.0 Storage"),
        Paragraph(text="Store finished product in a controlled warehouse area."),
    ]
    new_paragraphs = [
        Paragraph(text="1.0 Purpose"),
        Paragraph(text="This SOP defines tablet compression for all product lines."),
        Paragraph(text="2.0 Storage"),
        Paragraph(text="Store finished product in a controlled warehouse area."),
        Paragraph(text="Compression force shall be maintained at 18 kN."),
    ]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.txt", "new.txt")

    change_types = [c.change_type for c in result.changes]
    assert "moved_paragraph" in change_types
    assert "numeric_change" in change_types


def test_moved_paragraph_content_change_is_filed_under_the_new_section(monkeypatch):
    def fail_if_called(unresolved):
        raise AssertionError("no AI call expected")

    monkeypatch.setattr(llm_classifier, "classify_changes_batch", fail_if_called)

    old_paragraphs = [
        Paragraph(text="1.0 Purpose"),
        Paragraph(text="This SOP defines tablet compression for all product lines."),
        Paragraph(text="Compression force shall be maintained at 15 kN."),
        Paragraph(text="2.0 Storage"),
        Paragraph(text="Store finished product in a controlled warehouse area."),
    ]
    new_paragraphs = [
        Paragraph(text="1.0 Purpose"),
        Paragraph(text="This SOP defines tablet compression for all product lines."),
        Paragraph(text="2.0 Storage"),
        Paragraph(text="Store finished product in a controlled warehouse area."),
        Paragraph(text="Compression force shall be maintained at 18 kN."),
    ]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.txt", "new.txt")

    numeric = [c for c in result.changes if c.change_type == "numeric_change"]
    assert len(numeric) == 1
    # Filed where the paragraph now lives, so it groups with that section's other
    # changes in the UI - not under the move row's compound "old -> new" label.
    assert numeric[0].section == "2.0 Storage"


def test_moved_paragraph_content_change_carries_its_own_real_risk(monkeypatch):
    def fail_if_called(unresolved):
        raise AssertionError("no AI call expected")

    monkeypatch.setattr(llm_classifier, "classify_changes_batch", fail_if_called)

    old_paragraphs = [
        Paragraph(text="1.0 Purpose"),
        Paragraph(text="This SOP defines tablet compression for all product lines."),
        Paragraph(text="Compression force shall be maintained at 15 kN."),
        Paragraph(text="2.0 Storage"),
        Paragraph(text="Store finished product in a controlled warehouse area."),
    ]
    new_paragraphs = [
        Paragraph(text="1.0 Purpose"),
        Paragraph(text="This SOP defines tablet compression for all product lines."),
        Paragraph(text="2.0 Storage"),
        Paragraph(text="Store finished product in a controlled warehouse area."),
        Paragraph(text="Compression force shall be maintained at 18 kN."),
    ]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.txt", "new.txt")

    numeric = [c for c in result.changes if c.change_type == "numeric_change"]
    # The whole point of the fix: severity rides on its own row instead of being
    # downgraded to the move row's default.
    assert numeric[0].ai_risk_level == "High"


def test_moved_table_content_that_was_also_edited_reports_both(monkeypatch):
    def fail_if_called(unresolved):
        raise AssertionError("no AI call expected")

    monkeypatch.setattr(llm_classifier, "classify_changes_batch", fail_if_called)

    position = TableCoordinate(table_id=0, row=1, col=2)
    old_paragraphs = [
        Paragraph(text="1.0 Purpose"),
        Paragraph(text="This SOP defines the tablet compression process fully."),
        Paragraph(text="Assay limit 95.0 percent", from_table=True, table_position=position),
        Paragraph(text="2.0 Storage"),
        Paragraph(text="Store finished product in a controlled warehouse area."),
    ]
    new_paragraphs = [
        Paragraph(text="1.0 Purpose"),
        Paragraph(text="This SOP defines the tablet compression process fully."),
        Paragraph(text="2.0 Storage"),
        Paragraph(text="Store finished product in a controlled warehouse area."),
        Paragraph(text="Assay limit 98.0 percent", from_table=True, table_position=position),
    ]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.txt", "new.txt")

    change_types = [c.change_type for c in result.changes]
    assert "moved_table_content" in change_types
    numeric = [c for c in result.changes if c.change_type == "numeric_change"]
    assert len(numeric) == 1
    assert numeric[0].source == "Table"
    assert numeric[0].new_table_position == position


def test_pure_move_with_identical_text_produces_only_the_move_row(monkeypatch):
    # Regression guard: the fix must not duplicate-report relocations that
    # didn't change anything.
    def fail_if_called(unresolved):
        raise AssertionError("no AI call expected for an unchanged relocation")

    monkeypatch.setattr(llm_classifier, "classify_changes_batch", fail_if_called)

    old_paragraphs = [
        Paragraph(text="1.0 Purpose"),
        Paragraph(text="This SOP defines tablet compression for all product lines."),
        Paragraph(text="Compression force shall be maintained at 15 kN."),
        Paragraph(text="2.0 Storage"),
        Paragraph(text="Store finished product in a controlled warehouse area."),
    ]
    new_paragraphs = [
        Paragraph(text="1.0 Purpose"),
        Paragraph(text="This SOP defines tablet compression for all product lines."),
        Paragraph(text="2.0 Storage"),
        Paragraph(text="Store finished product in a controlled warehouse area."),
        Paragraph(text="Compression force shall be maintained at 15 kN."),
    ]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.txt", "new.txt")

    assert len(result.changes) == 1
    assert result.changes[0].change_type == "moved_paragraph"


def test_moved_paragraph_with_semantic_only_change_is_ai_classified(monkeypatch):
    def fake_classify(unresolved):
        return [
            LLMClassification(
                change_id=item["change_id"],
                change_type="role_responsibility_change",
                reason="Approval responsibility changed.",
                confidence=0.9,
            )
            for item in unresolved
        ]

    monkeypatch.setattr(llm_classifier, "classify_changes_batch", fake_classify)

    old_paragraphs = [
        Paragraph(text="1.0 Purpose"),
        Paragraph(text="This SOP defines tablet compression for all product lines."),
        Paragraph(text="The QC Manager shall approve the completed batch record."),
        Paragraph(text="2.0 Storage"),
        Paragraph(text="Store finished product in a controlled warehouse area."),
    ]
    new_paragraphs = [
        Paragraph(text="1.0 Purpose"),
        Paragraph(text="This SOP defines tablet compression for all product lines."),
        Paragraph(text="2.0 Storage"),
        Paragraph(text="Store finished product in a controlled warehouse area."),
        Paragraph(text="The QA Manager shall approve the completed batch record."),
    ]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.txt", "new.txt")

    change_types = [c.change_type for c in result.changes]
    assert "moved_paragraph" in change_types
    assert "role_responsibility_change" in change_types
    # The internal placeholder must never survive into the report.
    assert "pending_llm_classification" not in change_types


def test_moved_paragraph_tells_the_ai_which_change_types_regex_already_caught(monkeypatch):
    # A moved paragraph carrying BOTH a numeric edit and a residual semantic edit
    # must pass the already_detected hint through, so the AI classifies only the
    # genuinely distinct change instead of restating the numeric one.
    captured = {}

    def capturing_classify(unresolved):
        captured["items"] = unresolved
        return [
            LLMClassification(
                change_id=item["change_id"],
                change_type="role_responsibility_change",
                reason="Approval responsibility changed.",
                confidence=0.9,
            )
            for item in unresolved
        ]

    monkeypatch.setattr(llm_classifier, "classify_changes_batch", capturing_classify)

    old_paragraphs = [
        Paragraph(text="1.0 Purpose"),
        Paragraph(text="This SOP defines tablet compression for all product lines."),
        Paragraph(text="The QC Manager shall verify 15 kN compression force."),
        Paragraph(text="2.0 Storage"),
        Paragraph(text="Store finished product in a controlled warehouse area."),
    ]
    new_paragraphs = [
        Paragraph(text="1.0 Purpose"),
        Paragraph(text="This SOP defines tablet compression for all product lines."),
        Paragraph(text="2.0 Storage"),
        Paragraph(text="Store finished product in a controlled warehouse area."),
        Paragraph(text="The QA Manager shall verify 18 kN compression force."),
    ]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.txt", "new.txt")

    assert captured["items"], "the AI classifier should have been called"
    assert captured["items"][0]["already_detected"] == ["numeric_change"]
    change_types = [c.change_type for c in result.changes]
    assert "numeric_change" in change_types
    assert "role_responsibility_change" in change_types


def test_real_document_move_with_edit_reports_both_move_and_content_change(monkeypatch):
    # Real-document regression guard for this branch's whole reason to exist.
    # Every other test for this fix uses hand-built Paragraph objects, which skip
    # extraction, sectioning, and real section matching - so nothing else would
    # notice if a drift in those made this paragraph stop registering as a move
    # and the content change went silently missing again.
    #
    # In C1, "The stability chamber log..." moves from 8.0 Final Disposition into
    # the brand-new section 10.0 Long-Term Sample Retention AND is reworded.
    def fake_classify(unresolved):
        return [
            LLMClassification(
                change_id=item["change_id"],
                change_type="clarification_no_meaning_change",
                reason="Wording expanded for clarity.",
                confidence=0.9,
            )
            for item in unresolved
        ]

    monkeypatch.setattr(llm_classifier, "classify_changes_batch", fake_classify)

    old_paragraphs = extract_text(os.path.join(_TEST_DOCS_DIR, "C1_v1.docx"), "docx")
    new_paragraphs = extract_text(os.path.join(_TEST_DOCS_DIR, "C1_v2.docx"), "docx")

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "C1_v1.docx", "C1_v2.docx")

    old_text = "The stability chamber log shall be filed with the batch record."
    new_text = (
        "The stability chamber log shall be filed together with the batch record "
        "for long-term retention."
    )

    moves = [
        c for c in result.changes
        if c.change_type == "moved_paragraph" and c.old_text == old_text
    ]
    assert len(moves) == 1
    assert moves[0].section == "8.0 Final Disposition -> 10.0 Long-Term Sample Retention"
    assert moves[0].new_text == new_text

    # The edit that the move used to hide: reported separately, filed under the
    # section the paragraph now lives in.
    content_changes = [
        c for c in result.changes
        if c.change_type != "moved_paragraph"
        and c.old_text == old_text
        and c.new_text == new_text
    ]
    assert len(content_changes) == 1
    assert content_changes[0].section == "10.0 Long-Term Sample Retention"


def test_pipeline_reports_insertion_driven_renumbering_as_cascading(monkeypatch):
    # One insertion renumbers everything below it. Those rows are mechanical
    # consequences of a single edit and must be distinguishable from a
    # deliberate renumbering so a reviewer can filter them out.
    def fail_if_called(unresolved):
        raise AssertionError("no AI call expected for heading-only changes")

    monkeypatch.setattr(llm_classifier, "classify_changes_batch", fail_if_called)

    old_paragraphs = [
        Paragraph(text="1.0 Purpose"),
        Paragraph(text="This SOP defines the tablet compression process."),
        Paragraph(text="2.0 Equipment"),
        Paragraph(text="A rotary tablet press with standard tooling is used."),
        Paragraph(text="3.0 Records"),
        Paragraph(text="Batch records are retained for six years."),
    ]
    new_paragraphs = [
        Paragraph(text="1.0 Purpose"),
        Paragraph(text="This SOP defines the tablet compression process."),
        Paragraph(text="2.0 Safety"),
        Paragraph(text="Operators shall wear eye protection at all times."),
        Paragraph(text="3.0 Equipment"),
        Paragraph(text="A rotary tablet press with standard tooling is used."),
        Paragraph(text="4.0 Records"),
        Paragraph(text="Batch records are retained for six years."),
    ]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.txt", "new.txt")

    change_types = [c.change_type for c in result.changes]
    assert change_types.count("section_renumbered_cascade") == 2
    assert "section_renumbered" not in change_types
    assert "section_added" in change_types


def test_pipeline_reports_an_added_table_row_once_instead_of_per_cell(monkeypatch):
    # The whole point of the feature: one structural edit, one row in the report.
    def fail_if_called(unresolved):
        raise AssertionError("no AI call expected")

    monkeypatch.setattr(llm_classifier, "classify_changes_batch", fail_if_called)

    position = lambda r, c: TableCoordinate(table_id=0, row=r, col=c)
    old_paragraphs = [
        Paragraph(text="4.0 Procedure"),
        Paragraph(text="Parameter", from_table=True, table_position=position(0, 0)),
        Paragraph(text="Target", from_table=True, table_position=position(0, 1)),
        Paragraph(text="Compression Force", from_table=True, table_position=position(1, 0)),
        Paragraph(text="15 kN", from_table=True, table_position=position(1, 1)),
    ]
    new_paragraphs = [
        Paragraph(text="4.0 Procedure"),
        Paragraph(text="Parameter", from_table=True, table_position=position(0, 0)),
        Paragraph(text="Target", from_table=True, table_position=position(0, 1)),
        Paragraph(text="Compression Force", from_table=True, table_position=position(1, 0)),
        Paragraph(text="15 kN", from_table=True, table_position=position(1, 1)),
        Paragraph(text="Hardness", from_table=True, table_position=position(2, 0)),
        Paragraph(text="8 kg", from_table=True, table_position=position(2, 1)),
    ]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "old.txt", "new.txt")

    change_types = [c.change_type for c in result.changes]
    assert change_types.count("table_row_added") == 1
    # The two cells of the added row must not also be reported individually.
    assert "added_table_content" not in change_types


def test_pipeline_a_text_box_whose_anchor_changed_reports_no_heading_change(monkeypatch):
    # Adding a section above a text box shifts its anchor without the box moving.
    # This happens for real in the corpus (7.0 References -> 8.0 Training Requirements)
    # and must stay silent: anchoring is traceability, not a finding.
    monkeypatch.setattr(llm_classifier, "classify_changes_batch", lambda items, *a, **k: [])

    old_paragraphs = [
        Paragraph(text="1.0 Scope", paragraph_index=0, is_heading=True),
        Paragraph(text="Applies to all batches.", paragraph_index=1),
        Paragraph(text="Text Box 1 (1.0 Scope, after paragraph 1)", paragraph_index=2, is_heading=True),
        Paragraph(text="CAUTION: verify calibration.", paragraph_index=3),
    ]
    new_paragraphs = [
        Paragraph(text="1.0 Scope", paragraph_index=0, is_heading=True),
        Paragraph(text="Applies to all batches.", paragraph_index=1),
        Paragraph(text="Text Box 1 (2.0 Responsibilities, after paragraph 1)", paragraph_index=2, is_heading=True),
        Paragraph(text="CAUTION: verify calibration.", paragraph_index=3),
    ]

    result = pipeline.compare_documents(old_paragraphs, new_paragraphs, "v1.docx", "v2.docx")

    assert [c for c in result.changes if c.change_type == "section_heading_changed"] == []
