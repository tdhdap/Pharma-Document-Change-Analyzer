import zipfile

import fitz  # PyMuPDF, used here only to build test fixtures
from docx import Document as DocxDocument
from docx.shared import Inches, Pt
from lxml import etree
from docx.oxml import parse_xml
from docx.oxml.ns import qn

from app.extraction import extract_text, _docx_header_footer_specs, _header_footer_heading_text
from app.extraction import _iter_text_box_paragraphs
from app.extraction import _footnotes_root, _footnote_content_by_id
from app.extraction import _anchor_label
from app.sectioning import split_into_sections


_DRAWINGML_TEXTBOX_XML = """<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
    xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
    xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
  <wp:inline>
    <wp:extent cx="1828800" cy="1143000"/>
    <a:graphic>
      <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
        <wps:wsp>
          <wps:txbx>
            <w:txbxContent>{paragraphs}</w:txbxContent>
          </wps:txbx>
        </wps:wsp>
      </a:graphicData>
    </a:graphic>
  </wp:inline>
</w:drawing>"""

_VML_TEXTBOX_XML = """<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    xmlns:v="urn:schemas-microsoft-com:vml"
    xmlns:o="urn:schemas-microsoft-com:office:office">
  <v:shape style="width:150pt;height:80pt">
    <v:textbox>
      <w:txbxContent>{paragraphs}</w:txbxContent>
    </v:textbox>
  </v:shape>
</w:pict>"""


def _paragraph_xml(*run_texts):
    runs = "".join(f'<w:r><w:t xml:space="preserve">{t}</w:t></w:r>' for t in run_texts)
    return f"<w:p>{runs}</w:p>"


def _add_text_box(container_element, paragraph_texts, vml=False):
    """Inject a real text box (DrawingML by default, VML if vml=True) containing the
    given paragraphs into container_element (e.g. doc.element.body, a table cell's
    _tc, or a header/footer's _element - any lxml element works generically).
    python-docx has no API to add a text box, so this constructs the raw OOXML
    directly - the exact technique verified empirically before writing this plan."""
    p = container_element.makeelement(qn("w:p"), {})
    r = p.makeelement(qn("w:r"), {})
    p.append(r)
    paragraphs_xml = "".join(_paragraph_xml(text) for text in paragraph_texts)
    template = _VML_TEXTBOX_XML if vml else _DRAWINGML_TEXTBOX_XML
    drawing = etree.fromstring(template.format(paragraphs=paragraphs_xml).encode())
    r.append(drawing)
    container_element.append(p)


def _add_footnote_reference(paragraph, footnote_id):
    """Inject a real <w:footnoteReference w:id="..."/> into the given python-docx
    Paragraph's XML - python-docx has no API to add a footnote reference, so this
    constructs the raw OOXML directly, the same technique already used for text boxes."""
    p = paragraph._p
    r = p.makeelement(qn("w:r"), {})
    ref = r.makeelement(qn("w:footnoteReference"), {qn("w:id"): str(footnote_id)})
    r.append(ref)
    p.append(r)


def _save_docx_with_footnotes(doc, path, footnotes):
    """Save doc, then splice a real word/footnotes.xml part into the resulting .docx
    package (content-type override + relationship + the part itself). python-docx has
    no API to add this part, so this manipulates the OPC zip package directly - the
    exact technique verified empirically before writing this plan. `footnotes` is a
    list of (footnote_id, paragraph_texts) tuples for the real (non-boilerplate)
    footnotes; the two Word-internal separator footnotes are always included. Callers
    must have already used _add_footnote_reference for each corresponding id before
    calling doc.save() - this only splices the footnotes.xml part and its relationship,
    not the body's own <w:footnoteReference> elements."""
    doc.save(path)
    footnote_blocks = "".join(
        f'<w:footnote w:id="{fid}">' + "".join(_paragraph_xml(t) for t in texts) + "</w:footnote>"
        for fid, texts in footnotes
    )
    footnotes_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="0">'
        '<w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
        f'{footnote_blocks}'
        '</w:footnotes>'
    )
    with zipfile.ZipFile(path, "r") as zin:
        names = zin.namelist()
        content_types_xml = zin.read("[Content_Types].xml").decode("utf-8")
        rels_xml = zin.read("word/_rels/document.xml.rels").decode("utf-8")
        other = {n: zin.read(n) for n in names if n not in ("[Content_Types].xml", "word/_rels/document.xml.rels")}

    content_types_xml = content_types_xml.replace(
        "</Types>",
        '<Override PartName="/word/footnotes.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>',
    )
    rels_xml = rels_xml.replace(
        "</Relationships>",
        '<Relationship Id="rIdFootnotes" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" '
        'Target="footnotes.xml"/></Relationships>',
    )

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        zout.writestr("[Content_Types].xml", content_types_xml)
        zout.writestr("word/_rels/document.xml.rels", rels_xml)
        zout.writestr("word/footnotes.xml", footnotes_xml)
        for name, data in other.items():
            zout.writestr(name, data)


def test_extract_txt_splits_on_blank_lines(tmp_path):
    file_path = tmp_path / "doc.txt"
    file_path.write_text("First paragraph.\n\nSecond paragraph.\n\nThird paragraph.")

    paragraphs = extract_text(str(file_path), "txt")

    assert [p.text for p in paragraphs] == [
        "First paragraph.", "Second paragraph.", "Third paragraph.",
    ]
    assert paragraphs[0].paragraph_index == 0
    assert paragraphs[2].paragraph_index == 2
    assert all(p.page is None for p in paragraphs)
    assert all(p.is_heading is False for p in paragraphs)


def test_extract_docx_reads_paragraphs(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Heading One")
    doc.add_paragraph("")  # blank paragraphs should be skipped
    doc.add_paragraph("Body text.")
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")

    assert [p.text for p in paragraphs] == ["Heading One", "Body text."]
    assert paragraphs[0].paragraph_index == 0
    assert paragraphs[1].paragraph_index == 1


def test_extract_docx_tags_heading_style_paragraphs(tmp_path):
    """Isolates the structural "Heading 1" style signal: the run's font size is pinned
    to the body baseline (11pt) so the font-size signal cannot also fire here, ensuring
    this test actually discriminates the HEADING_STYLE_PREFIXES check from the size check."""
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    para = doc.add_paragraph(style="Heading 1")
    run = para.add_run("5.2 Sample Preparation")
    run.font.size = Pt(11)
    doc.add_paragraph("Weigh 10 mg of sample and dilute to volume.")
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")

    assert paragraphs[0].text == "5.2 Sample Preparation"
    assert paragraphs[0].is_heading is True
    assert paragraphs[1].is_heading is False


def test_extract_pdf_tags_page_numbers(tmp_path):
    file_path = tmp_path / "doc.pdf"
    pdf = fitz.open()
    page1 = pdf.new_page()
    page1.insert_text((72, 72), "Page one text.")
    page2 = pdf.new_page()
    page2.insert_text((72, 72), "Page two text.")
    pdf.save(str(file_path))
    pdf.close()

    paragraphs = extract_text(str(file_path), "pdf")

    assert any(p.page == 1 and "Page one text." in p.text for p in paragraphs)
    assert any(p.page == 2 and "Page two text." in p.text for p in paragraphs)
    assert all(p.is_heading is False for p in paragraphs)


def test_extract_pdf_tags_toc_entries_as_headings(tmp_path):
    file_path = tmp_path / "doc.pdf"
    pdf = fitz.open()
    page1 = pdf.new_page()
    page1.insert_text((72, 72), "5.2 Sample Preparation")
    page1.insert_text((72, 100), "10 mg of sample was weighed and diluted.")
    pdf.set_toc([[1, "5.2 Sample Preparation", 1]])
    pdf.save(str(file_path))
    pdf.close()

    paragraphs = extract_text(str(file_path), "pdf")

    heading_paragraphs = [p for p in paragraphs if p.is_heading]
    assert len(heading_paragraphs) == 1
    assert heading_paragraphs[0].text == "5.2 Sample Preparation"


def test_extract_pdf_preserves_wrapped_paragraphs_without_toc(tmp_path):
    """Regression test: wrapped paragraphs (with internal newlines) should remain as single Paragraphs
    when they don't match a TOC entry. This prevents spurious line-level diffs when paragraphs
    re-flow differently between document versions (margin/font changes)."""
    file_path = tmp_path / "doc.pdf"
    pdf = fitz.open()
    page1 = pdf.new_page()
    # Insert a multi-line block with no TOC entry: should be kept as one paragraph
    page1.insert_text((72, 72), "This is a long paragraph that\nwraps across multiple lines\nwithout any TOC entry.")
    pdf.save(str(file_path))
    pdf.close()

    paragraphs = extract_text(str(file_path), "pdf")

    # Should be exactly one paragraph (not fragmented into 3 lines)
    assert len(paragraphs) == 1
    assert paragraphs[0].page == 1
    assert paragraphs[0].is_heading is False
    assert "wraps across multiple lines" in paragraphs[0].text


def test_extract_pdf_tags_font_size_only_heading(tmp_path):
    file_path = tmp_path / "doc.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "1.0 Scope", fontsize=11)
    page.insert_text((72, 100), "This procedure applies to all testing in the QC lab.", fontsize=11)
    page.insert_text((72, 140), "Sample Preparation", fontsize=18)
    page.insert_text((72, 165), "Weigh 10 mg of sample and dilute to volume with mobile phase.", fontsize=11)
    pdf.save(str(file_path))
    pdf.close()

    paragraphs = extract_text(str(file_path), "pdf")

    # "1.0 Scope" is deliberately the same 11pt size as the body text around it (a
    # heading-shaped distractor at body size); only "Sample Preparation" at 18pt
    # clears the baseline+2.0pt threshold and should be flagged.
    heading_texts = {p.text for p in paragraphs if p.is_heading}
    assert heading_texts == {"Sample Preparation"}


def test_extract_pdf_does_not_flag_normal_body_text_as_heading(tmp_path):
    file_path = tmp_path / "doc.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "This is a completely normal sentence with no special formatting.", fontsize=11)
    pdf.save(str(file_path))
    pdf.close()

    paragraphs = extract_text(str(file_path), "pdf")

    assert all(not p.is_heading for p in paragraphs)


def test_extract_pdf_does_not_flag_size_below_threshold_as_heading(tmp_path):
    """A short, heading-shaped line whose size is only +1.0pt above the body baseline
    (below the +2.0pt threshold) must not be flagged as a heading."""
    file_path = tmp_path / "doc.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Weigh 10 mg of sample and dilute to volume with mobile phase.", fontsize=11)
    page.insert_text((72, 100), "Sample Preparation", fontsize=12)
    page.insert_text((72, 130), "Inject into the HPLC system for analysis.", fontsize=11)
    pdf.save(str(file_path))
    pdf.close()

    paragraphs = extract_text(str(file_path), "pdf")

    assert all(not p.is_heading for p in paragraphs)


def test_extract_pdf_does_not_flag_large_font_long_text_as_heading(tmp_path):
    """A large-font block whose text fails the heading shape guard (a long sentence
    ending in a period) must not be flagged as a heading, even though its size clears
    the threshold."""
    file_path = tmp_path / "doc.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Weigh 10 mg of sample and dilute to volume with mobile phase.", fontsize=11)
    page.insert_text(
        (72, 110),
        "This is a long sentence with a large font size but it ends with a period.",
        fontsize=18,
    )
    page.insert_text((72, 150), "Inject into the HPLC system for analysis.", fontsize=11)
    pdf.save(str(file_path))
    pdf.close()

    paragraphs = extract_text(str(file_path), "pdf")

    assert all(not p.is_heading for p in paragraphs)


def test_extract_text_rejects_unknown_type(tmp_path):
    file_path = tmp_path / "doc.xyz"
    file_path.write_text("content")

    try:
        extract_text(str(file_path), "xyz")
        assert False, "expected ValueError"
    except ValueError:
        pass


def test_extract_docx_tags_font_size_only_heading(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Weigh 10 mg of sample and dilute to volume.")
    p = doc.add_paragraph()
    run = p.add_run("Sample Preparation")
    run.bold = True
    run.font.size = Pt(16)
    doc.add_paragraph("Inject into the HPLC system.")
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")

    heading_paragraphs = [p for p in paragraphs if p.is_heading]
    assert len(heading_paragraphs) == 1
    assert heading_paragraphs[0].text == "Sample Preparation"


def test_extract_docx_does_not_flag_size_below_threshold_as_heading(tmp_path):
    """Test that a paragraph with explicit font size just below the heading threshold
    (11.0pt baseline + 2.0pt delta = 13.0pt) is NOT flagged as heading, even if text looks
    like a heading."""
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    p = doc.add_paragraph()
    run = p.add_run("Sample Preparation")
    run.font.size = Pt(12)  # Only +1.0pt above baseline, below the +2.0pt threshold
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")

    assert len(paragraphs) == 1
    assert paragraphs[0].text == "Sample Preparation"
    assert paragraphs[0].is_heading is False  # Should NOT be flagged (size 12 < threshold 13)


def test_extract_docx_uses_resolved_normal_style_size_as_baseline(tmp_path):
    """When the document's "Normal" style resolves a concrete font size, the baseline
    must be that resolved size (not the hardcoded 11.0pt fallback). Here Normal is set
    to 14pt, so +1pt (15pt) is below the +2.0pt threshold and must NOT be flagged, while
    +2pt (16pt) is at the threshold and MUST be flagged."""
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.styles["Normal"].font.size = Pt(14)

    p_below = doc.add_paragraph()
    p_below.add_run("Sample Preparation").font.size = Pt(15)

    p_at = doc.add_paragraph()
    p_at.add_run("Sample Analysis").font.size = Pt(16)

    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")

    by_text = {p.text: p.is_heading for p in paragraphs}
    assert by_text["Sample Preparation"] is False
    assert by_text["Sample Analysis"] is True


def test_extract_docx_does_not_flag_large_font_long_text_as_heading(tmp_path):
    """Test that a paragraph with large font size (above threshold) but whose text
    fails the shape guard (long sentence with period) is NOT flagged as heading."""
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    p = doc.add_paragraph()
    run = p.add_run("This is a long sentence with large font size but it ends with a period.")
    run.font.size = Pt(16)  # Clearly above threshold (11.0 + 2.0 = 13.0)
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")

    assert len(paragraphs) == 1
    assert paragraphs[0].is_heading is False  # Should NOT be flagged (shape guard fails due to period)


def test_extract_docx_extracts_table_cell_content(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("5.0 Acceptance Criteria", style="Heading 1")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Parameter"
    table.cell(0, 1).text = "Limit"
    table.cell(1, 0).text = "Assay"
    table.cell(1, 1).text = "95.0% to 105.0%"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert "Parameter" in texts
    assert "Limit" in texts
    assert "Assay" in texts
    assert "95.0% to 105.0%" in texts


def test_extract_docx_extracts_nested_table_cell_content(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    table = doc.add_table(rows=1, cols=1)
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested cell text"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert "Nested cell text" in texts


def test_extract_docx_detects_word_style_heading_inside_table_cell(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.paragraphs[0].text = "5.2 Sample Preparation"
    cell.paragraphs[0].style = doc.styles["Heading 1"]
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    heading_paragraphs = [p for p in paragraphs if p.is_heading]

    assert len(heading_paragraphs) == 1
    assert heading_paragraphs[0].text == "5.2 Sample Preparation"


def test_extract_docx_detects_font_size_heading_inside_table_cell(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Weigh 10 mg of sample and dilute to volume.")
    table = doc.add_table(rows=1, cols=1)
    p = table.cell(0, 0).paragraphs[0]
    run = p.add_run("Sample Preparation")
    run.bold = True
    run.font.size = Pt(16)
    doc.add_paragraph("Inject into the HPLC system.")
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    heading_paragraphs = [p for p in paragraphs if p.is_heading]

    assert len(heading_paragraphs) == 1
    assert heading_paragraphs[0].text == "Sample Preparation"


def test_extract_docx_without_tables_headers_or_footers_is_unchanged(tmp_path):
    """Regression guard: this refactor must not change extraction for documents
    that don't use any of the new features."""
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("1.0 Scope", style="Heading 1")
    doc.add_paragraph("This procedure applies to all testing.")
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")

    assert len(paragraphs) == 2
    assert paragraphs[0].text == "1.0 Scope"
    assert paragraphs[0].is_heading is True
    assert paragraphs[1].text == "This procedure applies to all testing."
    assert paragraphs[1].is_heading is False


def test_extract_docx_adds_page_header_section_when_header_is_set(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Body content here.")
    doc.sections[0].header.paragraphs[0].text = "Confidential SOP-1234"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert "Page Header" in texts
    header_index = texts.index("Page Header")
    assert paragraphs[header_index].is_heading is True
    assert "Confidential SOP-1234" in texts
    assert "Page Footer" not in texts


def test_extract_docx_adds_page_footer_section_when_footer_is_set(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Body content here.")
    doc.sections[0].footer.paragraphs[0].text = "Page footer text"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert "Page Footer" in texts
    footer_index = texts.index("Page Footer")
    assert paragraphs[footer_index].is_heading is True
    assert "Page footer text" in texts
    assert "Page Header" not in texts


def test_extract_docx_omits_page_header_and_footer_when_neither_is_set(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Body content here.")
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert "Page Header" not in texts
    assert "Page Footer" not in texts


def test_extract_docx_page_header_and_footer_both_present_appear_in_order(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Body content here.")
    doc.sections[0].header.paragraphs[0].text = "Header text"
    doc.sections[0].footer.paragraphs[0].text = "Footer text"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert texts.index("Page Header") < texts.index("Header text")
    assert texts.index("Header text") < texts.index("Page Footer")
    assert texts.index("Page Footer") < texts.index("Footer text")


def test_extract_docx_horizontally_merged_cell_extracted_once(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "MERGED HEADER CELL"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "B"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert texts.count("MERGED HEADER CELL") == 1
    assert texts.count("A") == 1
    assert texts.count("B") == 1


def test_extract_docx_vertically_merged_cell_extracted_once(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0))
    table.cell(0, 0).text = "VMERGED"
    table.cell(0, 1).text = "TOP RIGHT"
    table.cell(1, 1).text = "BOTTOM RIGHT"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert texts.count("VMERGED") == 1
    assert texts.count("TOP RIGHT") == 1
    assert texts.count("BOTTOM RIGHT") == 1


def test_extract_docx_whitespace_only_header_produces_no_page_header_section(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Body content here.")
    doc.sections[0].header.paragraphs[0].text = "   "
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert "Page Header" not in texts


def test_extract_docx_table_inside_header_is_extracted(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Body content here.")
    header = doc.sections[0].header
    header.paragraphs[0].text = "SOP-1234"
    table = header.add_table(rows=1, cols=2, width=Inches(6))
    table.cell(0, 0).text = "Header cell A"
    table.cell(0, 1).text = "Header cell B"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert "Header cell A" in texts
    assert "Header cell B" in texts


def test_extract_docx_multiple_sections_get_distinct_header_headings(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("First section body.")
    doc.sections[0].header.paragraphs[0].text = "First Header"
    doc.add_section()
    # A newly added section's header defaults to is_linked_to_previous=True, in which
    # case its "paragraphs" proxy the previous section's header paragraphs (so writing
    # through it would silently overwrite "First Header" instead of creating a second,
    # distinct header). Explicitly unlink it first to get a genuinely separate header.
    doc.sections[1].header.is_linked_to_previous = False
    doc.sections[1].header.paragraphs[0].text = "Second Header"
    doc.add_paragraph("Second section body.")
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert "Page Header (Section 1)" in texts
    assert "Page Header (Section 2)" in texts
    assert "Page Header" not in texts  # unqualified name must not appear once there are 2 sections

    section1_index = texts.index("Page Header (Section 1)")
    section2_index = texts.index("Page Header (Section 2)")
    assert texts[section1_index + 1] == "First Header"
    assert texts[section2_index + 1] == "Second Header"


def test_extract_docx_table_only_document_with_no_body_paragraphs(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Only cell A"
    table.cell(0, 1).text = "Only cell B"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert "Only cell A" in texts
    assert "Only cell B" in texts


def test_extract_docx_all_caps_table_cell_stays_in_its_section(tmp_path):
    """Test that ALL-CAPS table cell text doesn't fragment into its own heading section.
    The key is that allow_text_pattern_heading=False for table cells, so the sectioning
    logic won't detect "HPLC" as a heading."""
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("2.0 Acceptance Criteria", style="Heading 1")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Method"
    table.cell(0, 1).text = "HPLC"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    sections = split_into_sections(paragraphs)

    # Should have exactly 1 section (the main one under "2.0 Acceptance Criteria")
    # "HPLC" should NOT become its own section heading
    assert len(sections) == 1
    assert sections[0].heading == "2.0 Acceptance Criteria"
    body_texts = [p.text for p in sections[0].paragraphs]
    assert "HPLC" in body_texts


def test_extract_docx_all_caps_header_text_stays_in_page_header_section(tmp_path):
    """Test that ALL-CAPS header text doesn't fragment into its own heading.
    The key is that allow_text_pattern_heading=False for header paragraphs."""
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Body content here.")
    doc.sections[0].header.paragraphs[0].text = "CONFIDENTIAL"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    sections = split_into_sections(paragraphs)

    # Should have 2 sections: Preamble (body) and Page Header
    # "CONFIDENTIAL" should NOT become its own section heading under Page Header
    assert len(sections) == 2
    assert sections[0].heading == "Preamble"
    assert sections[1].heading == "Page Header"
    header_body = [p.text for p in sections[1].paragraphs]
    assert "CONFIDENTIAL" in header_body


def test_extract_docx_font_size_heading_inside_table_cell_still_detected(tmp_path):
    """Regression guard: structural signals (font-size) must still work
    inside table cells even though the text-pattern fallback is now
    disallowed there."""
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Weigh 10 mg of sample and dilute to volume.")
    table = doc.add_table(rows=1, cols=1)
    p = table.cell(0, 0).paragraphs[0]
    run = p.add_run("Sample Preparation")
    run.bold = True
    run.font.size = Pt(16)
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    sections = split_into_sections(paragraphs)

    # Font-size should still work inside table cells, so "Sample Preparation"
    # should be detected as a heading and create its own section
    assert len(sections) == 2
    assert sections[0].heading == "Preamble"
    assert sections[1].heading == "Sample Preparation"


def test_extract_docx_all_caps_body_paragraph_still_a_heading(tmp_path):
    """Regression guard: body-paragraph ALL-CAPS detection (not from a
    table or header/footer) is completely unaffected by this change."""
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("SCOPE")
    doc.add_paragraph("This procedure applies to all lab testing.")
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    sections = split_into_sections(paragraphs)

    # Body paragraphs should still have allow_text_pattern_heading=True,
    # so "SCOPE" should be detected as a heading
    assert len(sections) == 1
    assert sections[0].heading == "SCOPE"
    body_texts = [p.text for p in sections[0].paragraphs]
    assert "This procedure applies to all lab testing." in body_texts


def test_extract_docx_table_cell_has_from_table_true(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("2.0 Acceptance Criteria", style="Heading 1")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Method"
    table.cell(0, 1).text = "HPLC"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    by_text = {p.text: p for p in paragraphs}

    assert by_text["2.0 Acceptance Criteria"].from_table is False
    assert by_text["Method"].from_table is True
    assert by_text["HPLC"].from_table is True


def test_extract_docx_body_paragraph_has_from_table_false(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("1.0 Scope", style="Heading 1")
    doc.add_paragraph("This procedure applies to all lab testing.")
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")

    assert all(p.from_table is False for p in paragraphs)


def test_extract_docx_header_paragraph_not_from_table(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Body content here.")
    doc.sections[0].header.paragraphs[0].text = "Confidential"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    header_para = next(p for p in paragraphs if p.text == "Confidential")

    assert header_para.from_table is False


def test_extract_docx_table_inside_header_has_from_table_true(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Body content here.")
    header = doc.sections[0].header
    header.paragraphs[0].text = "SOP-1234"
    table = header.add_table(rows=1, cols=1, width=Inches(6))
    table.cell(0, 0).text = "Header Table Cell"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    cell_para = next(p for p in paragraphs if p.text == "Header Table Cell")

    assert cell_para.from_table is True


def test_extract_docx_table_cell_gets_correct_grid_coordinate(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "TopLeft"
    table.cell(0, 1).text = "TopRight"
    table.cell(1, 0).text = "BottomLeft"
    table.cell(1, 1).text = "BottomRight"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    by_text = {p.text: p for p in paragraphs}

    assert by_text["TopLeft"].table_position.row == 0
    assert by_text["TopLeft"].table_position.col == 0
    assert by_text["TopRight"].table_position.row == 0
    assert by_text["TopRight"].table_position.col == 1
    assert by_text["BottomLeft"].table_position.row == 1
    assert by_text["BottomLeft"].table_position.col == 0
    assert by_text["BottomRight"].table_position.row == 1
    assert by_text["BottomRight"].table_position.col == 1
    # All four cells belong to the same (only) table in the document.
    table_ids = {p.table_position.table_id for p in by_text.values()}
    assert len(table_ids) == 1


def test_extract_docx_non_table_paragraph_has_no_table_position(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("1.0 Scope", style="Heading 1")
    doc.add_paragraph("This procedure applies to all lab testing.")
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")

    assert all(p.table_position is None for p in paragraphs)


def test_extract_docx_horizontally_merged_cell_anchors_at_top_left(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "MERGED HEADER CELL"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "B"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    by_text = {p.text: p for p in paragraphs}

    assert by_text["MERGED HEADER CELL"].table_position.row == 0
    assert by_text["MERGED HEADER CELL"].table_position.col == 0
    assert by_text["A"].table_position.row == 1
    assert by_text["A"].table_position.col == 0
    assert by_text["B"].table_position.row == 1
    assert by_text["B"].table_position.col == 1


def test_extract_docx_vertically_merged_cell_anchors_at_top_left(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0))
    table.cell(0, 0).text = "VMERGED"
    table.cell(0, 1).text = "TOP RIGHT"
    table.cell(1, 1).text = "BOTTOM RIGHT"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    by_text = {p.text: p for p in paragraphs}

    assert by_text["VMERGED"].table_position.row == 0
    assert by_text["VMERGED"].table_position.col == 0
    assert by_text["TOP RIGHT"].table_position.row == 0
    assert by_text["TOP RIGHT"].table_position.col == 1
    assert by_text["BOTTOM RIGHT"].table_position.row == 1
    assert by_text["BOTTOM RIGHT"].table_position.col == 1


def test_extract_docx_multiple_body_tables_get_distinct_sequential_ids(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    first_table = doc.add_table(rows=1, cols=1)
    first_table.cell(0, 0).text = "First Table Cell"
    doc.add_paragraph("Some text between the two tables.")
    second_table = doc.add_table(rows=1, cols=1)
    second_table.cell(0, 0).text = "Second Table Cell"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    by_text = {p.text: p for p in paragraphs}

    first_id = by_text["First Table Cell"].table_position.table_id
    second_id = by_text["Second Table Cell"].table_position.table_id
    assert first_id != second_id
    assert second_id > first_id


def test_extract_docx_body_and_header_tables_share_the_global_id_counter(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    body_table = doc.add_table(rows=1, cols=1)
    body_table.cell(0, 0).text = "Body Table Cell"
    doc.add_paragraph("Body content here.")
    header = doc.sections[0].header
    header.paragraphs[0].text = "SOP-1234"
    header_table = header.add_table(rows=1, cols=1, width=Inches(6))
    header_table.cell(0, 0).text = "Header Table Cell"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    by_text = {p.text: p for p in paragraphs}

    body_id = by_text["Body Table Cell"].table_position.table_id
    header_id = by_text["Header Table Cell"].table_position.table_id
    assert body_id != header_id


def test_extract_docx_nested_table_gets_its_own_table_id_and_coordinates(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    outer_table = doc.add_table(rows=1, cols=2)
    outer_table.cell(0, 0).text = "Outer Cell"
    nested_table = outer_table.cell(0, 1).add_table(rows=1, cols=1)
    nested_table.cell(0, 0).text = "Nested Cell"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    by_text = {p.text: p for p in paragraphs}

    outer_position = by_text["Outer Cell"].table_position
    nested_position = by_text["Nested Cell"].table_position

    assert outer_position.table_id != nested_position.table_id
    # The nested table's own single cell is at its own grid position (0, 0),
    # not the outer cell's position (0, 1) that contains it.
    assert nested_position.row == 0
    assert nested_position.col == 0
    assert outer_position.row == 0
    assert outer_position.col == 0


def test_specs_include_default_header_and_footer_when_set():
    doc = DocxDocument()
    doc.sections[0].header.paragraphs[0].text = "Header text"
    doc.sections[0].footer.paragraphs[0].text = "Footer text"

    specs = _docx_header_footer_specs(doc)

    assert len(specs) == 2
    header_spec = next(s for s in specs if s[0] == "header")
    footer_spec = next(s for s in specs if s[0] == "footer")
    assert (header_spec[1], header_spec[2]) == (0, "")
    assert header_spec[3].paragraphs[0].text == "Header text"
    assert (footer_spec[1], footer_spec[2]) == (0, "")
    assert footer_spec[3].paragraphs[0].text == "Footer text"


def test_specs_omit_default_header_and_footer_when_unset():
    doc = DocxDocument()

    specs = _docx_header_footer_specs(doc)

    assert specs == []


def test_specs_omit_first_page_variant_when_toggle_is_off():
    doc = DocxDocument()
    doc.sections[0].first_page_header.paragraphs[0].text = "First page header text"
    # different_first_page_header_footer deliberately left False (default) - Word
    # would never actually display this content.

    specs = _docx_header_footer_specs(doc)

    assert not any(s[2] == "First Page" for s in specs)


def test_specs_include_first_page_variant_when_toggle_is_on():
    doc = DocxDocument()
    doc.sections[0].different_first_page_header_footer = True
    doc.sections[0].first_page_header.paragraphs[0].text = "First page header text"

    specs = _docx_header_footer_specs(doc)

    first_page_specs = [s for s in specs if s[2] == "First Page"]
    assert len(first_page_specs) == 1
    assert first_page_specs[0][0] == "header"
    assert first_page_specs[0][1] == 0
    assert first_page_specs[0][3].paragraphs[0].text == "First page header text"


def test_specs_omit_even_page_variant_when_document_toggle_is_off():
    doc = DocxDocument()
    doc.sections[0].even_page_header.paragraphs[0].text = "Even page header text"
    # doc.settings.odd_and_even_pages_header_footer deliberately left False (default).

    specs = _docx_header_footer_specs(doc)

    assert not any(s[2] == "Even Page" for s in specs)


def test_specs_include_even_page_variant_when_document_toggle_is_on():
    doc = DocxDocument()
    doc.settings.odd_and_even_pages_header_footer = True
    doc.sections[0].even_page_header.paragraphs[0].text = "Even page header text"

    specs = _docx_header_footer_specs(doc)

    even_page_specs = [s for s in specs if s[2] == "Even Page"]
    assert len(even_page_specs) == 1
    assert even_page_specs[0][0] == "header"
    assert even_page_specs[0][3].paragraphs[0].text == "Even page header text"


def test_specs_are_ordered_all_headers_then_all_footers_across_sections():
    doc = DocxDocument()
    doc.sections[0].header.paragraphs[0].text = "Section 1 header"
    doc.sections[0].footer.paragraphs[0].text = "Section 1 footer"
    doc.add_section()
    doc.sections[1].header.is_linked_to_previous = False
    doc.sections[1].header.paragraphs[0].text = "Section 2 header"
    doc.sections[1].footer.is_linked_to_previous = False
    doc.sections[1].footer.paragraphs[0].text = "Section 2 footer"

    specs = _docx_header_footer_specs(doc)

    kinds_and_sections = [(s[0], s[1]) for s in specs]
    assert kinds_and_sections == [("header", 0), ("header", 1), ("footer", 0), ("footer", 1)]


def test_heading_text_default_single_section():
    assert _header_footer_heading_text("header", 0, "", multi_section=False) == "Page Header"
    assert _header_footer_heading_text("footer", 0, "", multi_section=False) == "Page Footer"


def test_heading_text_default_multi_section():
    assert _header_footer_heading_text("header", 1, "", multi_section=True) == "Page Header (Section 2)"


def test_heading_text_first_page_single_section():
    assert _header_footer_heading_text("header", 0, "First Page", multi_section=False) == "Page Header (First Page)"


def test_heading_text_first_page_multi_section():
    result = _header_footer_heading_text("footer", 1, "First Page", multi_section=True)
    assert result == "Page Footer (Section 2, First Page)"


def test_heading_text_even_page_single_section():
    assert _header_footer_heading_text("header", 0, "Even Page", multi_section=False) == "Page Header (Even Page)"


def test_heading_text_even_page_multi_section():
    result = _header_footer_heading_text("header", 1, "Even Page", multi_section=True)
    assert result == "Page Header (Section 2, Even Page)"


def test_extract_docx_first_page_header_extracted_when_toggle_is_on(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Body content here.")
    doc.sections[0].different_first_page_header_footer = True
    doc.sections[0].first_page_header.paragraphs[0].text = "First page header text"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert "Page Header (First Page)" in texts
    assert "First page header text" in texts


def test_extract_docx_first_page_header_omitted_when_toggle_is_off(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Body content here.")
    doc.sections[0].first_page_header.paragraphs[0].text = "First page header text"
    # different_first_page_header_footer deliberately left False (default) -
    # Word would never actually display this content.
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert "First page header text" not in texts
    assert not any("First Page" in t for t in texts)


def test_extract_docx_even_page_header_extracted_when_document_toggle_is_on(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Body content here.")
    doc.settings.odd_and_even_pages_header_footer = True
    doc.sections[0].even_page_header.paragraphs[0].text = "Even page header text"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert "Page Header (Even Page)" in texts
    assert "Even page header text" in texts


def test_extract_docx_even_page_header_omitted_when_document_toggle_is_off(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Body content here.")
    doc.sections[0].even_page_header.paragraphs[0].text = "Even page header text"
    # doc.settings.odd_and_even_pages_header_footer deliberately left False (default).
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert "Even page header text" not in texts
    assert not any("Even Page" in t for t in texts)


def test_extract_docx_first_page_header_in_multi_section_document_is_fully_qualified(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("First section body.")
    doc.add_section()
    doc.sections[1].different_first_page_header_footer = True
    # As with a regular header (see the distinct-headers test above), a newly added
    # section's first_page_header defaults to is_linked_to_previous=True. Because section
    # 1 never enabled different_first_page_header_footer, python-docx's inheritance walk
    # would otherwise resolve the write all the way back to section 1's (unused) first-page
    # definition instead of creating one for section 2. Explicitly unlink first.
    doc.sections[1].first_page_header.is_linked_to_previous = False
    doc.sections[1].first_page_header.paragraphs[0].text = "Section 2 first page header"
    doc.add_paragraph("Second section body.")
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert "Page Header (Section 2, First Page)" in texts
    assert "Section 2 first page header" in texts


def test_extract_docx_first_page_header_inherited_when_linked_even_with_toggle_on(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("First section body.")
    doc.sections[0].different_first_page_header_footer = True
    doc.sections[0].first_page_header.paragraphs[0].text = "Section 1 first page header"
    doc.add_section()
    doc.sections[1].different_first_page_header_footer = True
    # Section 2's first_page_header is left linked to previous (default) - it should
    # inherit section 1's content rather than getting its own pseudo-section.
    doc.add_paragraph("Second section body.")
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert "Page Header (Section 1, First Page)" in texts
    assert "Page Header (Section 2, First Page)" not in texts


def test_extract_docx_table_inside_first_page_header_gets_table_position(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Body content here.")
    doc.sections[0].different_first_page_header_footer = True
    first_page_header = doc.sections[0].first_page_header
    table = first_page_header.add_table(rows=1, cols=1, width=Inches(6))
    table.cell(0, 0).text = "First page header table cell"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    cell_para = next(p for p in paragraphs if p.text == "First page header table cell")

    assert cell_para.from_table is True
    assert cell_para.table_position is not None


def test_iter_text_box_paragraphs_finds_drawingml_text_box():
    doc = DocxDocument()
    _add_text_box(doc.element.body, ["Text box paragraph one.", "Text box paragraph two."])

    groups = [group for _txbx, group in _iter_text_box_paragraphs(doc.element.body, doc)]

    assert len(groups) == 1
    assert [p.text for p in groups[0]] == ["Text box paragraph one.", "Text box paragraph two."]


def test_iter_text_box_paragraphs_finds_vml_text_box():
    doc = DocxDocument()
    _add_text_box(doc.element.body, ["Legacy VML text box paragraph."], vml=True)

    groups = [group for _txbx, group in _iter_text_box_paragraphs(doc.element.body, doc)]

    assert len(groups) == 1
    assert [p.text for p in groups[0]] == ["Legacy VML text box paragraph."]


def test_iter_text_box_paragraphs_joins_multiple_runs():
    doc = DocxDocument()
    p = doc.element.body.makeelement(qn("w:p"), {})
    r = p.makeelement(qn("w:r"), {})
    p.append(r)
    paragraphs_xml = _paragraph_xml("Multi-", "run", " sentence.")
    drawing = etree.fromstring(_DRAWINGML_TEXTBOX_XML.format(paragraphs=paragraphs_xml).encode())
    r.append(drawing)
    doc.element.body.append(p)

    groups = [group for _txbx, group in _iter_text_box_paragraphs(doc.element.body, doc)]

    assert len(groups) == 1
    assert [p.text for p in groups[0]] == ["Multi-run sentence."]


def test_iter_text_box_paragraphs_returns_nothing_when_none_present():
    doc = DocxDocument()
    doc.add_paragraph("Ordinary paragraph, no text box.")

    groups = [group for _txbx, group in _iter_text_box_paragraphs(doc.element.body, doc)]

    assert groups == []


def test_iter_text_box_paragraphs_finds_multiple_text_boxes_as_separate_groups():
    doc = DocxDocument()
    _add_text_box(doc.element.body, ["First box text."])
    _add_text_box(doc.element.body, ["Second box text."])

    groups = [group for _txbx, group in _iter_text_box_paragraphs(doc.element.body, doc)]

    assert len(groups) == 2
    assert [p.text for p in groups[0]] == ["First box text."]
    assert [p.text for p in groups[1]] == ["Second box text."]


def test_iter_text_box_paragraphs_finds_text_box_inside_table_cell():
    doc = DocxDocument()
    table = doc.add_table(rows=1, cols=1)
    cell_element = table.cell(0, 0)._tc
    _add_text_box(cell_element, ["Table cell text box."])

    groups = [group for _txbx, group in _iter_text_box_paragraphs(doc.element.body, doc)]

    assert len(groups) == 1
    assert [p.text for p in groups[0]] == ["Table cell text box."]


def test_iter_text_box_paragraphs_handles_nested_text_box_as_separate_group():
    doc = DocxDocument()
    p = doc.element.body.makeelement(qn("w:p"), {})
    r = p.makeelement(qn("w:r"), {})
    p.append(r)
    inner_paragraphs_xml = _paragraph_xml("Inner nested box text.")
    inner_drawing_xml = _DRAWINGML_TEXTBOX_XML.format(paragraphs=inner_paragraphs_xml)
    outer_paragraphs_xml = (
        '<w:p><w:r><w:t xml:space="preserve">Outer box own text.</w:t></w:r>'
        f'<w:r>{inner_drawing_xml}</w:r></w:p>'
    )
    outer_drawing = etree.fromstring(_DRAWINGML_TEXTBOX_XML.format(paragraphs=outer_paragraphs_xml).encode())
    r.append(outer_drawing)
    doc.element.body.append(p)

    groups = [group for _txbx, group in _iter_text_box_paragraphs(doc.element.body, doc)]

    assert len(groups) == 2
    assert [p.text for p in groups[0]] == ["Outer box own text."]
    assert [p.text for p in groups[1]] == ["Inner nested box text."]


def test_extract_docx_text_box_in_body_is_extracted(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Normal body paragraph.")
    _add_text_box(doc.element.body, ["Text box content here."])
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert "Text Box 1 (Preamble, after paragraph 1)" in texts
    text_box_index = texts.index("Text Box 1 (Preamble, after paragraph 1)")
    assert paragraphs[text_box_index].is_heading is True
    assert "Text box content here." in texts


def test_extract_docx_text_box_content_has_correct_default_fields(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    _add_text_box(doc.element.body, ["Text box content here."])
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    content_para = next(p for p in paragraphs if p.text == "Text box content here.")

    assert content_para.from_table is False
    assert content_para.table_position is None


def test_extract_docx_multiple_text_boxes_get_sequential_numbers(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    _add_text_box(doc.element.body, ["First box."])
    _add_text_box(doc.element.body, ["Second box."])
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert "Text Box 1 (Preamble, at start)" in texts
    assert "Text Box 2 (Preamble, at start)" in texts
    assert texts.index("Text Box 1 (Preamble, at start)") < texts.index("First box.")
    assert texts.index("Text Box 2 (Preamble, at start)") < texts.index("Second box.")


def test_extract_docx_empty_text_box_produces_no_pseudo_section(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Normal body paragraph.")
    _add_text_box(doc.element.body, [""])
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert not any(t.startswith("Text Box ") for t in texts)


def test_extract_docx_text_box_inside_active_first_page_header_is_extracted(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Body content here.")
    doc.sections[0].different_first_page_header_footer = True
    first_page_header_element = doc.sections[0].first_page_header._element
    _add_text_box(first_page_header_element, ["First page header text box."])
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert "Text Box 1 (Page Header (First Page))" in texts
    assert "First page header text box." in texts


def test_extract_docx_text_box_inside_inactive_variant_is_not_extracted(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Body content here.")
    # different_first_page_header_footer deliberately left False - Word would never
    # display this content, and _docx_header_footer_specs correctly excludes this
    # variant entirely regardless of the text box injected into it.
    first_page_header_element = doc.sections[0].first_page_header._element
    _add_text_box(first_page_header_element, ["Inactive first page header text box."])
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    texts = [p.text for p in paragraphs]

    assert "Inactive first page header text box." not in texts
    assert not any(t.startswith("Text Box ") for t in texts)


def test_extract_docx_without_text_boxes_is_unchanged(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("1.0 Scope", style="Heading 1")
    doc.add_paragraph("This procedure applies to all testing.")
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")

    assert len(paragraphs) == 2
    assert paragraphs[0].text == "1.0 Scope"
    assert paragraphs[1].text == "This procedure applies to all testing."


_MC_ALTERNATE_CONTENT_TEXTBOX_XML = """<mc:AlternateContent xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
    xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
    xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
    xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
    xmlns:v="urn:schemas-microsoft-com:vml"
    xmlns:o="urn:schemas-microsoft-com:office:office">
  <mc:Choice Requires="wps">
    <w:drawing>
      <wp:inline>
        <a:graphic>
          <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
            <wps:wsp>
              <wps:txbx>
                <w:txbxContent>{paragraphs}</w:txbxContent>
              </wps:txbx>
            </wps:wsp>
          </a:graphicData>
        </a:graphic>
      </wp:inline>
    </w:drawing>
  </mc:Choice>
  <mc:Fallback>
    <w:pict>
      <v:shape>
        <v:textbox>
          <w:txbxContent>{paragraphs}</w:txbxContent>
        </v:textbox>
      </v:shape>
    </w:pict>
  </mc:Fallback>
</mc:AlternateContent>"""


def _add_mc_alternate_content_text_box(container_element, paragraph_texts):
    """Inject a text box the way real Word 2010+ writes it: one
    mc:AlternateContent element wrapping BOTH a DrawingML (mc:Choice) and VML
    (mc:Fallback) copy of the identical content - verified empirically to match
    Word's actual output shape. Mirrors _add_text_box's construction pattern."""
    p = container_element.makeelement(qn("w:p"), {})
    r = p.makeelement(qn("w:r"), {})
    p.append(r)
    paragraphs_xml = "".join(_paragraph_xml(text) for text in paragraph_texts)
    mc = etree.fromstring(_MC_ALTERNATE_CONTENT_TEXTBOX_XML.format(paragraphs=paragraphs_xml).encode())
    r.append(mc)
    container_element.append(p)


def test_iter_text_box_paragraphs_deduplicates_word_mc_alternate_content():
    doc = DocxDocument()
    body = doc.element.body
    _add_mc_alternate_content_text_box(body, ["Caution: wear gloves."])

    groups = [group for _txbx, group in _iter_text_box_paragraphs(body, doc)]

    assert len(groups) == 1
    assert [para.text for para in groups[0]] == ["Caution: wear gloves."]


_DRAWINGML_TEXTBOX_WITH_TABLE_XML = """<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
    xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
    xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
  <wp:inline>
    <a:graphic>
      <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
        <wps:wsp>
          <wps:txbx>
            <w:txbxContent>
              <w:p><w:r><w:t xml:space="preserve">Caption above table.</w:t></w:r></w:p>
              <w:tbl>
                <w:tblPr/>
                <w:tblGrid><w:gridCol/><w:gridCol/></w:tblGrid>
                <w:tr>
                  <w:tc><w:p><w:r><w:t xml:space="preserve">Spec limit</w:t></w:r></w:p></w:tc>
                  <w:tc><w:p><w:r><w:t xml:space="preserve">NMT 2.0%</w:t></w:r></w:p></w:tc>
                </w:tr>
              </w:tbl>
            </w:txbxContent>
          </wps:txbx>
        </wps:wsp>
      </a:graphicData>
    </a:graphic>
  </wp:inline>
</w:drawing>"""


def test_iter_text_box_paragraphs_extracts_nested_table_content():
    doc = DocxDocument()
    body = doc.element.body
    p = body.makeelement(qn("w:p"), {})
    r = p.makeelement(qn("w:r"), {})
    p.append(r)
    drawing = etree.fromstring(_DRAWINGML_TEXTBOX_WITH_TABLE_XML.encode())
    r.append(drawing)
    body.append(p)

    groups = [group for _txbx, group in _iter_text_box_paragraphs(body, doc)]

    assert len(groups) == 1
    assert [para.text for para in groups[0]] == ["Caption above table.", "Spec limit", "NMT 2.0%"]


def test_footnotes_root_returns_none_when_document_has_no_footnotes():
    doc = DocxDocument()
    doc.add_paragraph("Plain paragraph, no footnotes.")

    assert _footnotes_root(doc) is None


def test_footnote_content_by_id_excludes_boilerplate_and_keys_by_id():
    footnotes_xml = (
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="0">'
        '<w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
        '<w:footnote w:id="1">' + _paragraph_xml("See ICH Q1A(R2) for stability testing requirements.") + '</w:footnote>'
        '</w:footnotes>'
    )
    root = parse_xml(footnotes_xml.encode())
    doc = DocxDocument()

    content_by_id = _footnote_content_by_id(root, doc)

    assert list(content_by_id.keys()) == ["1"]
    assert [para.text for para in content_by_id["1"]] == ["See ICH Q1A(R2) for stability testing requirements."]


def test_footnote_content_by_id_extracts_multi_paragraph_footnote():
    footnotes_xml = (
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:footnote w:id="1">'
        + _paragraph_xml("First line of the footnote.")
        + _paragraph_xml("Second line of the footnote.")
        + '</w:footnote>'
        '</w:footnotes>'
    )
    root = parse_xml(footnotes_xml.encode())
    doc = DocxDocument()

    content_by_id = _footnote_content_by_id(root, doc)

    assert [para.text for para in content_by_id["1"]] == [
        "First line of the footnote.", "Second line of the footnote.",
    ]


def test_footnote_content_by_id_extracts_nested_table_content():
    footnotes_xml = (
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:footnote w:id="1">'
        '<w:p><w:r><w:t xml:space="preserve">Reference standards:</w:t></w:r></w:p>'
        '<w:tbl>'
        '<w:tblPr/><w:tblGrid><w:gridCol/><w:gridCol/></w:tblGrid>'
        '<w:tr>'
        '<w:tc><w:p><w:r><w:t xml:space="preserve">USP</w:t></w:r></w:p></w:tc>'
        '<w:tc><w:p><w:r><w:t xml:space="preserve">Chapter 621</w:t></w:r></w:p></w:tc>'
        '</w:tr>'
        '</w:tbl>'
        '</w:footnote>'
        '</w:footnotes>'
    )
    root = parse_xml(footnotes_xml.encode())
    doc = DocxDocument()

    content_by_id = _footnote_content_by_id(root, doc)

    assert [para.text for para in content_by_id["1"]] == ["Reference standards:", "USP", "Chapter 621"]


def test_footnote_content_by_id_handles_multiple_non_contiguous_ids():
    footnotes_xml = (
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:footnote w:id="5">' + _paragraph_xml("First citation text.") + '</w:footnote>'
        '<w:footnote w:id="7">' + _paragraph_xml("Second citation text.") + '</w:footnote>'
        '</w:footnotes>'
    )
    root = parse_xml(footnotes_xml.encode())
    doc = DocxDocument()

    content_by_id = _footnote_content_by_id(root, doc)

    assert set(content_by_id.keys()) == {"5", "7"}
    assert [para.text for para in content_by_id["5"]] == ["First citation text."]
    assert [para.text for para in content_by_id["7"]] == ["Second citation text."]


def test_extract_docx_single_footnote_is_extracted(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    p = doc.add_paragraph("This is a claim with a footnote reference.")
    _add_footnote_reference(p, "1")
    _save_docx_with_footnotes(doc, str(file_path), [("1", ["See ICH Q1A(R2) for stability testing requirements."])])

    paragraphs = extract_text(str(file_path), "docx")

    texts = [par.text for par in paragraphs]
    assert texts == [
        "This is a claim with a footnote reference.",
        "Footnote 1 (Preamble, paragraph 1)",
        "See ICH Q1A(R2) for stability testing requirements.",
    ]
    heading_flags = {par.text: par.is_heading for par in paragraphs}
    assert heading_flags["Footnote 1 (Preamble, paragraph 1)"] is True
    assert heading_flags["See ICH Q1A(R2) for stability testing requirements."] is False


def test_extract_docx_multiple_footnotes_get_sequential_numbers_in_reference_order(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    p1 = doc.add_paragraph("First claim.")
    _add_footnote_reference(p1, "5")
    p2 = doc.add_paragraph("Second claim.")
    _add_footnote_reference(p2, "7")
    _save_docx_with_footnotes(doc, str(file_path), [
        ("5", ["First citation text."]),
        ("7", ["Second citation text."]),
    ])

    paragraphs = extract_text(str(file_path), "docx")

    texts = [par.text for par in paragraphs]
    assert texts == [
        "First claim.", "Second claim.",
        "Footnote 1 (Preamble, paragraph 1)", "First citation text.",
        "Footnote 2 (Preamble, paragraph 2)", "Second citation text.",
    ]


def test_extract_docx_footnote_content_has_correct_default_fields(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    p = doc.add_paragraph("Claim.")
    _add_footnote_reference(p, "1")
    _save_docx_with_footnotes(doc, str(file_path), [("1", ["Citation text."])])

    paragraphs = extract_text(str(file_path), "docx")

    footnote_para = next(par for par in paragraphs if par.text == "Citation text.")
    assert footnote_para.from_table is False
    assert footnote_para.table_position is None
    assert footnote_para.allow_text_pattern_heading is False


def test_extract_docx_footnote_with_nested_table_is_extracted(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    p = doc.add_paragraph("Claim needing a reference standard.")
    _add_footnote_reference(p, "1")
    doc.save(str(file_path))
    footnotes_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:footnote w:id="1">'
        '<w:p><w:r><w:t xml:space="preserve">Reference standards:</w:t></w:r></w:p>'
        '<w:tbl>'
        '<w:tblPr/><w:tblGrid><w:gridCol/><w:gridCol/></w:tblGrid>'
        '<w:tr>'
        '<w:tc><w:p><w:r><w:t xml:space="preserve">USP</w:t></w:r></w:p></w:tc>'
        '<w:tc><w:p><w:r><w:t xml:space="preserve">Chapter 621</w:t></w:r></w:p></w:tc>'
        '</w:tr>'
        '</w:tbl>'
        '</w:footnote>'
        '</w:footnotes>'
    )
    with zipfile.ZipFile(str(file_path), "r") as zin:
        names = zin.namelist()
        content_types_xml = zin.read("[Content_Types].xml").decode("utf-8")
        rels_xml = zin.read("word/_rels/document.xml.rels").decode("utf-8")
        other = {n: zin.read(n) for n in names if n not in ("[Content_Types].xml", "word/_rels/document.xml.rels")}
    content_types_xml = content_types_xml.replace(
        "</Types>",
        '<Override PartName="/word/footnotes.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>',
    )
    rels_xml = rels_xml.replace(
        "</Relationships>",
        '<Relationship Id="rIdFootnotes" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" '
        'Target="footnotes.xml"/></Relationships>',
    )
    with zipfile.ZipFile(str(file_path), "w", zipfile.ZIP_DEFLATED) as zout:
        zout.writestr("[Content_Types].xml", content_types_xml)
        zout.writestr("word/_rels/document.xml.rels", rels_xml)
        zout.writestr("word/footnotes.xml", footnotes_xml)
        for name, data in other.items():
            zout.writestr(name, data)

    paragraphs = extract_text(str(file_path), "docx")

    texts = [par.text for par in paragraphs]
    assert texts == [
        "Claim needing a reference standard.",
        "Footnote 1 (Preamble, paragraph 1)", "Reference standards:", "USP", "Chapter 621",
    ]
    table_para = next(par for par in paragraphs if par.text == "USP")
    assert table_para.from_table is False
    assert table_para.table_position is None


def test_extract_docx_separator_footnotes_are_never_extracted(tmp_path):
    # A real separator/continuationSeparator footnote's body has no actual text (just
    # a <w:separator/> marker element), and the body never actually references its id
    # either - so excluding it by w:type wouldn't visibly change round-trip output on
    # its own (nothing looks it up, and even if something did, it has no text to show).
    # To genuinely exercise the w:type gate end-to-end, this constructs a deliberately
    # unrealistic separator footnote that DOES carry real paragraph text, AND adds a
    # body reference to its id - proving it's excluded even when referenced, not just
    # coincidentally unreachable.
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    p = doc.add_paragraph("Claim.")
    _add_footnote_reference(p, "-1")
    _add_footnote_reference(p, "1")
    doc.save(str(file_path))
    footnotes_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:footnote w:type="separator" w:id="-1">'
        + _paragraph_xml("This text should never be extracted.") + '</w:footnote>'
        '<w:footnote w:id="1">' + _paragraph_xml("Real citation.") + '</w:footnote>'
        '</w:footnotes>'
    )
    with zipfile.ZipFile(str(file_path), "r") as zin:
        names = zin.namelist()
        content_types_xml = zin.read("[Content_Types].xml").decode("utf-8")
        rels_xml = zin.read("word/_rels/document.xml.rels").decode("utf-8")
        other = {n: zin.read(n) for n in names if n not in ("[Content_Types].xml", "word/_rels/document.xml.rels")}
    content_types_xml = content_types_xml.replace(
        "</Types>",
        '<Override PartName="/word/footnotes.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>',
    )
    rels_xml = rels_xml.replace(
        "</Relationships>",
        '<Relationship Id="rIdFootnotes" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" '
        'Target="footnotes.xml"/></Relationships>',
    )
    with zipfile.ZipFile(str(file_path), "w", zipfile.ZIP_DEFLATED) as zout:
        zout.writestr("[Content_Types].xml", content_types_xml)
        zout.writestr("word/_rels/document.xml.rels", rels_xml)
        zout.writestr("word/footnotes.xml", footnotes_xml)
        for name, data in other.items():
            zout.writestr(name, data)

    paragraphs = extract_text(str(file_path), "docx")

    texts = [par.text for par in paragraphs]
    assert "Real citation." in texts
    assert "This text should never be extracted." not in texts


def test_extract_docx_without_footnotes_is_unchanged(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Plain paragraph, no footnotes at all.")
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")

    assert [par.text for par in paragraphs] == ["Plain paragraph, no footnotes at all."]


def test_extract_docx_reference_to_missing_footnote_id_is_skipped(tmp_path):
    # Defensive case: a <w:footnoteReference> whose w:id has no matching <w:footnote>
    # in footnotes.xml (a malformed document) is skipped silently, and extraction of
    # everything else proceeds normally.
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    p = doc.add_paragraph("Claim with a dangling footnote reference.")
    _add_footnote_reference(p, "99")
    _save_docx_with_footnotes(doc, str(file_path), [("1", ["Unrelated real footnote, different id."])])

    paragraphs = extract_text(str(file_path), "docx")

    texts = [par.text for par in paragraphs]
    assert texts == ["Claim with a dangling footnote reference."]


def test_extract_docx_footnote_referenced_from_inside_text_box_is_extracted(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Ordinary body paragraph.")
    body = doc.element.body
    p = body.makeelement(qn("w:p"), {})
    r = p.makeelement(qn("w:r"), {})
    p.append(r)
    drawing_xml = """<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
    xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
    xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
  <wp:inline><a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
    <wps:wsp><wps:txbx>
      <w:txbxContent><w:p><w:r><w:t xml:space="preserve">Callout text with a footnote.</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r></w:p></w:txbxContent>
    </wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:inline>
</w:drawing>"""
    drawing = etree.fromstring(drawing_xml.encode())
    r.append(drawing)
    body.append(p)
    _save_docx_with_footnotes(doc, str(file_path), [("1", ["Footnote referenced from inside a text box."])])

    paragraphs = extract_text(str(file_path), "docx")

    texts = [par.text for par in paragraphs]
    assert texts == [
        "Ordinary body paragraph.",
        "Text Box 1 (Preamble, after paragraph 1)", "Callout text with a footnote.",
        "Footnote 1 (Preamble, after paragraph 1)", "Footnote referenced from inside a text box.",
    ]


def test_extract_docx_footnote_in_mc_alternate_content_text_box_is_extracted_once(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    body = doc.element.body
    p = body.makeelement(qn("w:p"), {})
    r = p.makeelement(qn("w:r"), {})
    p.append(r)
    mc_xml = """<mc:AlternateContent xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
    xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
    xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
    xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
    xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">
  <mc:Choice Requires="wps">
    <w:drawing><wp:inline><a:graphic><a:graphicData uri="x"><wps:wsp><wps:txbx>
      <w:txbxContent><w:p><w:r><w:t xml:space="preserve">Box text.</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r></w:p></w:txbxContent>
    </wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing>
  </mc:Choice>
  <mc:Fallback>
    <w:pict><v:shape><v:textbox>
      <w:txbxContent><w:p><w:r><w:t xml:space="preserve">Box text.</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r></w:p></w:txbxContent>
    </v:textbox></v:shape></w:pict>
  </mc:Fallback>
</mc:AlternateContent>"""
    mc = etree.fromstring(mc_xml.encode())
    r.append(mc)
    body.append(p)
    _save_docx_with_footnotes(doc, str(file_path), [("1", ["Should appear exactly once."])])

    paragraphs = extract_text(str(file_path), "docx")

    footnote_headings = [par.text for par in paragraphs if par.text.startswith("Footnote ")]
    assert footnote_headings == ["Footnote 1 (Preamble, at start)"]
    assert [par.text for par in paragraphs].count("Should appear exactly once.") == 1


def test_merge_spans_are_captured_for_both_directions(tmp_path):
    # row.cells returns a merged cell once per grid position it occupies, and all
    # those proxies wrap one w:tc element - so counting distinct row and column
    # indices per element gives both spans with one mechanism.
    file_path = tmp_path / "merged.docx"
    doc = DocxDocument()
    table = doc.add_table(rows=3, cols=3)
    for r in range(3):
        for c in range(3):
            table.rows[r].cells[c].text = f"r{r}c{c}"
    table.rows[0].cells[0].merge(table.rows[0].cells[1])   # horizontal, span 2
    table.rows[1].cells[2].merge(table.rows[2].cells[2])   # vertical, span 2
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    spans = {
        (p.table_position.row, p.table_position.col):
            (p.table_position.row_span, p.table_position.col_span)
        for p in paragraphs if p.table_position is not None
    }

    assert spans[(0, 0)] == (1, 2)
    assert spans[(1, 2)] == (2, 1)
    assert spans[(1, 0)] == (1, 1)
    assert spans[(2, 0)] == (1, 1)


def test_unmerged_table_cells_all_have_span_one(tmp_path):
    file_path = tmp_path / "plain.docx"
    doc = DocxDocument()
    table = doc.add_table(rows=2, cols=2)
    for r in range(2):
        for c in range(2):
            table.rows[r].cells[c].text = f"r{r}c{c}"
    doc.save(str(file_path))

    paragraphs = extract_text(str(file_path), "docx")
    positions = [p.table_position for p in paragraphs if p.table_position is not None]

    assert len(positions) == 4
    assert all(pos.row_span == 1 and pos.col_span == 1 for pos in positions)


def test_anchor_label_inside_an_emitted_paragraph():
    assert _anchor_label("4.0 Procedure", 0, True) == "4.0 Procedure, paragraph 1"
    assert _anchor_label("4.0 Procedure", 2, True) == "4.0 Procedure, paragraph 3"


def test_anchor_label_between_paragraphs_when_the_anchor_was_dropped():
    assert _anchor_label("8.0 Training", 1, False) == "8.0 Training, after paragraph 1"


def test_anchor_label_at_start_when_nothing_precedes_it():
    # "after paragraph 0" would be nonsense; a real corpus text box hits this.
    assert _anchor_label("9.0 Training Log", 0, False) == "9.0 Training Log, at start"


def test_extract_docx_text_box_anchor_names_the_enclosing_heading(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_heading("4.0 Procedure", level=1)
    doc.add_paragraph("Body paragraph under the heading.")
    _add_text_box(doc.element.body, ["Callout content."])
    doc.save(str(file_path))

    texts = [p.text for p in extract_text(str(file_path), "docx")]

    assert "Text Box 1 (4.0 Procedure, after paragraph 1)" in texts


def test_extract_docx_text_box_anchor_uses_a_text_pattern_heading(tmp_path):
    # The heading test must be sectioning._is_heading_paragraph, not model.is_heading.
    # An ALL-CAPS heading has no Heading style, so model.is_heading is False for it,
    # but split_into_sections treats it as a heading - 10 corpus documents rely on this.
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Opening body text.")
    doc.add_paragraph("RECORD RETENTION")
    doc.add_paragraph("Records are retained for five years.")
    _add_text_box(doc.element.body, ["Callout content."])
    doc.save(str(file_path))

    texts = [p.text for p in extract_text(str(file_path), "docx")]

    assert "Text Box 1 (RECORD RETENTION, after paragraph 1)" in texts


def test_extract_docx_two_text_boxes_at_one_position_stay_distinguishable(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    _add_text_box(doc.element.body, ["First box."])
    _add_text_box(doc.element.body, ["Second box."])
    doc.save(str(file_path))

    texts = [p.text for p in extract_text(str(file_path), "docx")]

    assert "Text Box 1 (Preamble, at start)" in texts
    assert "Text Box 2 (Preamble, at start)" in texts


def test_extract_docx_footnote_anchor_names_the_paragraph_it_sits_in(tmp_path):
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_heading("4.0 Procedure", level=1)
    doc.add_paragraph("First body paragraph.")
    p = doc.add_paragraph("Second body paragraph.")
    _add_footnote_reference(p, "1")
    _save_docx_with_footnotes(doc, str(file_path), [("1", ["Citation text."])])

    texts = [par.text for par in extract_text(str(file_path), "docx")]

    assert "Footnote 1 (4.0 Procedure, paragraph 2)" in texts


def test_extract_docx_footnote_in_a_paragraph_with_no_text_anchors_between(tmp_path):
    # A paragraph holding only the reference marker is dropped by
    # _docx_paragraph_to_model, so the shared rule reports it as between paragraphs.
    file_path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Body paragraph.")
    p = doc.add_paragraph("")
    _add_footnote_reference(p, "1")
    _save_docx_with_footnotes(doc, str(file_path), [("1", ["Citation text."])])

    texts = [par.text for par in extract_text(str(file_path), "docx")]

    assert "Footnote 1 (Preamble, after paragraph 1)" in texts
